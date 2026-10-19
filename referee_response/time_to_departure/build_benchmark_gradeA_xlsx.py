import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
import os
FE=os.environ.get('GRADEA_FE','fye1')       # 'fye1' (fiscal-year-end date FE) or 'fy' (fiscal-year FE) for the interaction sheet
R=pd.read_pickle('gradeA_res.pkl' if FE=='fye1' else 'gradeA_res_fy.pkl'); T3=R['T3']; INT=R['INT']; COR=R['COR']; YS=R['YS']; XB=R['XB']; XC=R['XC']; BY=R['BY']
PANELS=[('Panel A. Peer-Adjusted Returns',[y for y,_ in YS if not y.startswith('b_')]),('Panel B. Benchmark-Adjusted Returns',BY)]
TNR='Times New Roman'; SZ=10.5
F=Font(name=TNR,size=SZ); FB=Font(name=TNR,size=SZ,bold=True); FI=Font(name=TNR,size=SZ,italic=True); BLUE=Font(name=TNR,size=SZ,color='0000FF')
thin=Side(style='thin'); CENTER=Alignment(horizontal='center',vertical='center',wrap_text=True); RIGHT=Alignment(horizontal='right'); JUST=Alignment(horizontal='justify',wrap_text=True,vertical='top')
YL={'peer_adj_ret100':'All plan-years','peer_cs':'Plan-years with the grade-A benchmark-adjusted return','b_AB':'Reported + inferred grade A/B (as before)','b_A':'Reported + inferred grade A only','b_rep':'Reported benchmarks only'}
YD={'peer_adj_ret100':'Peer-adjusted','peer_cs':'Peer-adjusted','b_AB':'Benchmark-adjusted','b_A':'Benchmark-adjusted','b_rep':'Benchmark-adjusted'}
wb=Workbook(); E=wb.active; E.title='Estimates'
E['A1']='Estimates behind the tables. Blue = regression output (coefficient, SE, minimum cluster count). t and p are formulas; p uses the t distribution with (min clusters - 1) df. SE double clustered by plan and year, with the eigenvalue fix.'; E['A1'].font=FI
for j,h in enumerate(['Table','Benchmark','Column / X','Term','Coefficient','SE','Clusters (min)','t','p'],1): E.cell(row=3,column=j,value=h).font=FB
for c,w in zip('ABCD',[18,42,28,40]): E.column_dimensions[c].width=w
er=[4]
def reg(tab,y,col,term,r,G):
    i=er[0]; er[0]+=1
    for j,v in enumerate([tab,YD[y]+': '+YL[y],col,term],1): E.cell(row=i,column=j,value=v).font=F
    E.cell(row=i,column=5,value=r['est']).font=BLUE; E.cell(row=i,column=6,value=r['se']).font=BLUE; E.cell(row=i,column=7,value=G).font=BLUE
    E.cell(row=i,column=8,value=f'=IF(F{i}>0,E{i}/F{i},"")').font=F; E.cell(row=i,column=9,value=f'=IF(H{i}="","",TDIST(ABS(H{i}),G{i}-1,2))').font=F
    for j in [5,6,8,9]: E.cell(row=i,column=j).number_format='0.0000'
    return dict(coef=f'Estimates!E{i}',t=f'Estimates!H{i}',p=f'Estimates!I{i}')
def cf(ref): return f'=IF({ref["p"]}="","",TEXT({ref["coef"]},"0.000")&IF({ref["p"]}<0.01,"***",IF({ref["p"]}<0.05,"**",IF({ref["p"]}<0.1,"*",""))))'
def tf(ref): return f'=IF({ref["t"]}="","","("&TEXT({ref["t"]},"0.00")&")")'
def put(ws,r,c,ref,fnt=F):
    a=ws.cell(row=r,column=c,value=cf(ref)); a.font=fnt; a.alignment=RIGHT; b=ws.cell(row=r+1,column=c,value=tf(ref)); b.font=fnt; b.alignment=RIGHT
def note(ws,text,ncol,per=120):
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=1+ncol); ws['A2']=text; ws['A2'].font=F; ws['A2'].alignment=JUST; ws.row_dimensions[2].height=14*(len(text)//per+2)
BMNOTE=("Benchmark-Adjusted Return is the plan's return minus its policy benchmark return. 'Reported + inferred grade A/B' uses the benchmark return reported in the plan's financial report or, where none is stated, an inferred value graded A or B (the version in the earlier tables). "
        "'Reported + inferred grade A only' drops the inferred values graded B (28 plan-years, 10 of them conflicted). 'Reported benchmarks only' drops all inferred values. ")
# ---------------- Table III Panel B ----------------
ws=wb.create_sheet('Table III Panel B',0); ws.sheet_view.showGridLines=False; ws.column_dimensions['A'].width=34
for c in range(2,8): ws.column_dimensions[L(c)].width=13
ws['A1']='Table III: Peer- and Benchmark-Adjusted Returns by Benchmark Confidence'; ws['A1'].font=FB
note(ws,"This table re-estimates Table III. Panel A uses Peer-Adjusted Return, for all plan-years (as in the paper) and for the plan-years that also have the grade-A benchmark-adjusted return. Panel B uses Benchmark-Adjusted Return for three versions of the policy benchmark. "+BMNOTE+"Columns follow Table III: (1) base controls with year fixed effects; (2) all controls with year fixed effects; (3) plan fixed effects; (4) plan and year fixed effects; (5) and (6) code Conflicted CIO as 1 only in the last 5 or 10 years before the CIO departs. The standard errors are double clustered by plan and year. Coefficients marked with ***, **, and * are significant at the 1%, 5%, and 10% level, respectively.",6,105)
r=4
for c in range(6): x=ws.cell(row=r,column=2+c,value=f'({c+1})'); x.font=F; x.alignment=CENTER; x.border=Border(top=thin,bottom=thin)
ws.cell(row=r,column=1).border=Border(top=thin,bottom=thin); r+=1
T3ORDER=[(pt,y) for pt,ys in PANELS for y in ys]
for pt,y in T3ORDER:
    if y==PANELS[0][1][0] or y==PANELS[1][1][0]: ws.cell(row=r,column=1,value=pt).font=FB; ws.cell(row=r,column=1).border=Border(bottom=thin); r+=1
    ws.cell(row=r,column=1,value=YL[y]).font=FI; r+=1
    ws.cell(row=r,column=1,value='Conflicted CIO (0/1)').font=F
    for c,res in enumerate(T3[y]): put(ws,r,2+c,reg('Table III B',y,f'({c+1})','Conflicted CIO',res['conf'],res['G']))
    r+=2
    for lab,fn,fmt in [('Observations',lambda x:x['N'],'#,##0'),('Conflicted plan-years',lambda x:x['nconf'],'#,##0'),('Adjusted R-squared',lambda x:x['adjr2'],'0.000')]:
        ws.cell(row=r,column=1,value=lab).font=F
        for c,res in enumerate(T3[y]): x=ws.cell(row=r,column=2+c,value=fn(res)); x.font=F; x.alignment=RIGHT; x.number_format=fmt
        r+=1
    r+=1
for lab,vals in [('Controls',['Base']+['All']*5),('Plan FE',['No','No','Yes','Yes','No','No']),('Year FE',['Yes','Yes','No','Yes','Yes','Yes']),('Last n years',['','','','','5','10'])]:
    ws.cell(row=r,column=1,value=lab).font=F
    for c,v in enumerate(vals): x=ws.cell(row=r,column=2+c,value=v); x.font=F; x.alignment=CENTER
    r+=1
for c in range(1,8): ws.cell(row=r-1,column=c).border=Border(bottom=thin)
# ---------------- Interactions, benchmark-adjusted ----------------
wi=wb.create_sheet('Interactions Panel B',1); wi.sheet_view.showGridLines=False; wi.column_dimensions['A'].width=40
XS=XB+XC; NC=len(XS)
for c in range(2,2+NC): wi.column_dimensions[L(c)].width=13
wi['A1']='Interaction Specification: Peer- and Benchmark-Adjusted Returns by Benchmark Confidence'; wi['A1'].font=FB
note(wi,"This table repeats the interaction specification of xsec_interactions_fye1.xlsx (all Table III controls, "+('fiscal-year-end date fixed effects' if FE=='fye1' else 'fiscal-year fixed effects instead of fiscal-year-end date fixed effects')+", Conflicted CIO interacted with X). Panel A uses Peer-Adjusted Return, for all plan-years and for the plan-years that also have the grade-A benchmark-adjusted return; Panel B uses Benchmark-Adjusted Return for three versions of the policy benchmark. "+BMNOTE+"For binary X, Conflicted CIO is the effect when X = 0 and the italic row gives the effect when X = 1. Continuous X (log pay, relative pay filled within spell, standardized log CIO age) is centered at its regression-sample mean, so Conflicted CIO is the effect at average X, and the italic rows give the effect at the 25th and 75th percentiles. Log CIO Age (std.) is the paper's log_cio_age, ln(1 + age), standardized over the Table III sample (mean 3.950, standard deviation 0.174), so its interaction is the change in the effect of conflict for a one-standard-deviation increase in log age, about 10 years at the mean age of 52; that column drops the lagged log CIO age control. The standard errors are double clustered by plan and year. Coefficients marked with ***, **, and * are significant at the 1%, 5%, and 10% level, respectively.",NC,140)
r=4
for pt,y in T3ORDER:
    if y==PANELS[0][1][0] or y==PANELS[1][1][0]: wi.cell(row=r,column=1,value=pt).font=FB; r+=1
    wi.cell(row=r,column=1,value=YL[y]).font=FI; r+=1
    for j,(xv,xl) in enumerate(XS): x=wi.cell(row=r,column=2+j,value=f'X = {xl}'); x.font=F; x.alignment=CENTER; x.border=Border(top=thin,bottom=thin)
    wi.cell(row=r,column=1).border=Border(top=thin,bottom=thin); wi.row_dimensions[r].height=30; r+=1
    for term,tl,fnt in [('conf','Conflicted CIO (0/1)',F),('confX','Conflicted CIO x X',F),('X','X',F),('sum','Conflicted CIO effect when X = 1',FI),('p25','Conflicted CIO effect, 25th pct. of X',FI),('p75','Conflicted CIO effect, 75th pct. of X',FI)]:
        if term=='sum': r+=1
        wi.cell(row=r,column=1,value=tl).font=fnt
        for j,(xv,xl) in enumerate(XS):
            cont = xv in [a for a,_ in XC]
            if (term=='sum' and cont) or (term in ('p25','p75') and not cont): continue
            res=INT[y][xv]; put(wi,r,2+j,reg('Interactions B',y,xl,tl,res[term],res['G']),fnt)
        r+=2
    for lab,fn,fmt in [('Observations',lambda x:x['N'],'#,##0'),('Adjusted R-squared',lambda x:x['adjr2'],'0.000')]:
        wi.cell(row=r,column=1,value=lab).font=F
        for j,(xv,xl) in enumerate(XS): x=wi.cell(row=r,column=2+j,value=fn(INT[y][xv])); x.font=F; x.alignment=RIGHT; x.number_format=fmt
        r+=1
    for c in range(1,2+NC): wi.cell(row=r-1,column=c).border=Border(bottom=thin)
    r+=2
# ---------------- Correlations ----------------
wc=wb.create_sheet('Correlations',2); wc.sheet_view.showGridLines=False; wc.column_dimensions['A'].width=44
for c in range(2,9): wc.column_dimensions[L(c)].width=14
wc['A1']='Correlation of Benchmark-Adjusted and Peer-Adjusted Returns'; wc['A1'].font=FB
note(wc,"Correlations are computed over Table III sample plan-years (cio_dum = 1, fy <= 2020, curr1 = 0) where both returns are available. "+("'Within fiscal-year-end' demeans both returns by fiscal-year-end date (fye1)" if FE=='fye1' else "'Within fiscal year' demeans both returns by fiscal year (fy)")+" before correlating. "+BMNOTE,6,95)
r=4
hd=['Benchmark version','Plan-years','Pearson','Spearman',('Pearson, within fiscal-year-end' if FE=='fye1' else 'Pearson, within fiscal year'),'Pearson, conflicted plan-years','Pearson, other plan-years']
for j,h in enumerate(hd): x=wc.cell(row=r,column=1+j,value=h); x.font=F; x.alignment=CENTER; x.border=Border(top=thin,bottom=thin)
wc.row_dimensions[r].height=45; r+=1
for y in BY:
    c=COR[y]; vals=[YL[y],c['N'],c['pearson'],c['spearman'],c['within'],c['conf'],c['nonconf']]
    for j,v in enumerate(vals):
        x=wc.cell(row=r,column=1+j,value=float(v) if j>=2 else v); x.font=F if j==0 else BLUE; x.alignment=RIGHT if j else Alignment(horizontal='left'); x.number_format='0.000' if j>=2 else '#,##0'
    r+=1
for j in range(len(hd)): wc.cell(row=r-1,column=1+j).border=Border(bottom=thin)
for sh in (ws,wi,wc):
    sh.page_setup.orientation='landscape'; sh.sheet_properties.pageSetUpPr.fitToPage=True; sh.page_setup.fitToWidth=1; sh.page_setup.fitToHeight=0
out='/home/user/PensionData/referee_response/time_to_departure/'+('benchmark_gradeA_tables.xlsx' if FE=='fye1' else 'benchmark_gradeA_tables_fy.xlsx'); wb.save(out); print('saved',out)
