import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
R = pd.read_pickle('xsec_newbench_res_v2.pkl'); RES=R['RES']; YS=R['YS']; SPL={k:(sv,cl,labs,tab) for k,sv,cl,labs,tab in R['SPLITS']}
TNR='Times New Roman'; SZ=10.5
F=Font(name=TNR,size=SZ); FB=Font(name=TNR,size=SZ,bold=True); FI=Font(name=TNR,size=SZ,italic=True); BLUE=Font(name=TNR,size=SZ,color='0000FF')
thin=Side(style='thin'); CENTER=Alignment(horizontal='center',vertical='center',wrap_text=True); RIGHT=Alignment(horizontal='right'); JUST=Alignment(horizontal='justify',wrap_text=True,vertical='top')
wb=Workbook(); E=wb.active; E.title='Estimates'
E['A1']='Estimates behind the tables. Blue = regression output (coefficient, SE, minimum cluster count). t and p are formulas; p uses the t distribution with (min clusters - 1) df. SE double clustered by plan and fiscal year (plan and fiscal-year-end date for the financial-center split), with the eigenvalue fix.'; E['A1'].font=FI
for j,h in enumerate(['Split','Dependent variable','Term','Coefficient','SE','Clusters (min)','t','p'],1): E.cell(row=3,column=j,value=h).font=FB
E.column_dimensions['A'].width=12; E.column_dimensions['B'].width=48; E.column_dimensions['C'].width=24
er=[4]
def reg(split,yv,term,r,G):
    i=er[0]; er[0]+=1
    for j,v in enumerate([split,yv,term],1): E.cell(row=i,column=j,value=v).font=F
    E.cell(row=i,column=4,value=r['est']).font=BLUE; E.cell(row=i,column=5,value=r['se']).font=BLUE; E.cell(row=i,column=6,value=G).font=BLUE
    E.cell(row=i,column=7,value=f'=IF(E{i}>0,D{i}/E{i},"")').font=F; E.cell(row=i,column=8,value=f'=IF(G{i}="","",TDIST(ABS(G{i}),F{i}-1,2))').font=F
    for j in [4,5,7,8]: E.cell(row=i,column=j).number_format='0.0000'
    return dict(coef=f'Estimates!D{i}',t=f'Estimates!G{i}',p=f'Estimates!H{i}')
def cf(ref): return f'=IF({ref["p"]}="","n/a",TEXT({ref["coef"]},"0.000")&IF({ref["p"]}<0.01,"***",IF({ref["p"]}<0.05,"**",IF({ref["p"]}<0.1,"*",""))))'
def tf(ref): return f'=IF({ref["t"]}="","","("&TEXT({ref["t"]},"0.00")&")")'
def pf_(ref): return f'=IF({ref["p"]}="","",{ref["p"]})'
MAIN=[('peer_adj_ret100','Peer-Adjusted Return (t)'),('bench_adj_ret_pct','Benchmark-Adjusted Return, Original Benchmark (t)'),('b_AB','Benchmark-Adjusted Return, Updated Benchmark (t)')]
PAPER_FIRST_G0 = {'priv','fin'}   # the paper lists the 'not' group first in these panels
def panel(ws,r,key,ptitle):
    sv,cl,(l1,l0),tab = SPL[key]
    order = [('b',l0),('a',l1)] if key in PAPER_FIRST_G0 else [('a',l1),('b',l0)]
    ws.cell(row=r,column=1,value=ptitle).font=FB; r+=1
    for j,(yv,ylab) in enumerate(MAIN):
        c0=2+2*j; ws.merge_cells(start_row=r,start_column=c0,end_row=r,end_column=c0+1); c=ws.cell(row=r,column=c0,value=ylab); c.font=F; c.alignment=CENTER
        for cc in (c0,c0+1): ws.cell(row=r,column=cc).border=Border(top=thin)
    ws.cell(row=r,column=1).border=Border(top=thin); ws.row_dimensions[r].height=30; r+=1
    for j,(yv,_) in enumerate(MAIN):
        for k,(side,lab) in enumerate(order):
            c=ws.cell(row=r,column=2+2*j+k,value=lab); c.font=F; c.alignment=CENTER; c.border=Border(bottom=thin)
    ws.cell(row=r,column=1).border=Border(bottom=thin); ws.row_dimensions[r].height=30; r+=1
    ws.cell(row=r,column=1,value='Conflicted CIO').font=F
    for j,(yv,_) in enumerate(MAIN):
        x=RES[key][yv]
        for k,(side,lab) in enumerate(order):
            ref=reg(key,yv,f'{lab} x Conflicted',x[side],x['G'])
            a=ws.cell(row=r,column=2+2*j+k,value=cf(ref)); a.font=F; a.alignment=RIGHT; b=ws.cell(row=r+1,column=2+2*j+k,value=tf(ref)); b.font=F; b.alignment=RIGHT
    r+=3
    ws.cell(row=r,column=1,value='p-value: equal coefficients').font=F
    for j,(yv,_) in enumerate(MAIN):
        x=RES[key][yv]; ref=reg(key,yv,'difference',x['d'],x['G'])
        c0=2+2*j; ws.merge_cells(start_row=r,start_column=c0,end_row=r,end_column=c0+1); c=ws.cell(row=r,column=c0,value=pf_(ref)); c.font=F; c.alignment=CENTER; c.number_format='0.000'
    r+=1
    for lab,fn,fmt in [('Observations',lambda x:x['N'],'#,##0'),('Adjusted R-squared',lambda x:x['adjr2'],'0.000')]:
        ws.cell(row=r,column=1,value=lab).font=F
        for j,(yv,_) in enumerate(MAIN):
            c0=2+2*j; ws.merge_cells(start_row=r,start_column=c0,end_row=r,end_column=c0+1); c=ws.cell(row=r,column=c0,value=fn(RES[key][yv])); c.font=F; c.alignment=CENTER; c.number_format=fmt
        r+=1
    for lab in ['Controls','Year FE']:
        ws.cell(row=r,column=1,value=lab).font=F
        for j in range(3):
            c0=2+2*j; ws.merge_cells(start_row=r,start_column=c0,end_row=r,end_column=c0+1); c=ws.cell(row=r,column=c0,value='Yes'); c.font=F; c.alignment=CENTER
        r+=1
    for c in range(1,8): ws.cell(row=r-1,column=c).border=Border(bottom=thin)
    return r+2
common=("Each panel reports regressions of plan performance on Conflicted CIO interacted with the group indicator, with all control variables also interacted with the group indicator, so that the two coefficients are estimated in one regression. The first pair of columns uses Peer-Adjusted Return, as in the paper. The second pair uses Benchmark-Adjusted Return computed with the original policy benchmark, and the third pair uses Benchmark-Adjusted Return recomputed as the plan's return minus the updated policy benchmark "
 "(the benchmark return reported in the plan's financial report; where none is stated, a value inferred from the plan's multi-year benchmark returns or from its policy weights and component index returns, using only estimates graded A or B). The control variables are those of the corresponding regressions in the paper and are defined in Table I. All regressions include year fixed effects. The p-value row tests that the two Conflicted CIO coefficients are equal. ")
cl_txt = "The standard errors are double clustered by plan and year. Coefficients marked with ***, **, and * are significant at the 1%, 5%, and 10% level, respectively."
def sheet(name,title,intro,panels,cl_note=cl_txt):
    ws=wb.create_sheet(name); ws.sheet_view.showGridLines=False; ws.column_dimensions['A'].width=30
    for c in range(2,8): ws.column_dimensions[L(c)].width=14
    ws.cell(row=1,column=1,value=title).font=FB
    txt=intro+' '+common+cl_note
    ws.merge_cells('A2:G2'); ws['A2']=txt; ws['A2'].font=F; ws['A2'].alignment=JUST; ws.row_dimensions[2].height=15*(len(txt)//105+2)
    r=4
    for key,pt in panels: r=panel(ws,r,key,pt)
sheet('Table V','Table V: Conflicted CIOs: CIO Incentives and Plan Performance (updated benchmark)',
      "This table divides the sample by two proxies for a CIO's susceptibility to become conflicted: an indicator equal to 1 if the CIO is underpaid relative to the size of the plan, and an indicator equal to 1 for CIOs who are 46 years old or younger (missing when CIO age is missing).",
      [('underpaid','Panel A. Effect of Conflict for Underpaid CIOs'),('young','Panel B. Effect of Conflict for Younger CIOs')])
sheet('Table VI','Table VI: Conflicted CIOs: Alternative Pathways and Plan Performance (updated benchmark)',
      "This table divides the sample by two proxies for alternative pathways to the private sector: an indicator equal to 1 for CIOs with prior private firm experience, and an indicator equal to 1 for plans headquartered in financial centers.",
      [('priv','Panel A. Effect of Conflict for CIOs With and Without Prior Private Experience'),('fin','Panel B. Effect of Conflict for Plans Located in Financial Centers')],
      cl_note="The standard errors are double clustered by plan and year (by plan and fiscal-year-end date in Panel B, as in the paper's code). Coefficients marked with ***, **, and * are significant at the 1%, 5%, and 10% level, respectively.")
sheet('Table VII','Table VII: Conflicted CIOs: Governance and Plan Performance (updated benchmark)',
      "This table divides the sample by two proxies for the strength of plan governance: an indicator equal to 1 if the plan's discount rate is above 8%, and an indicator equal to 1 for plans with a funding ratio of at least 92% (missing when the funding ratio is missing).",
      [('rate','Panel A. Effect of Conflict Based on Liability Discount Rate'),('fund','Panel B. Effect of Conflict Based on Funding')])
# versions sheet
ws=wb.create_sheet('Benchmark versions'); ws.sheet_view.showGridLines=False; ws.column_dimensions['A'].width=60
for c in range(2,6): ws.column_dimensions[L(c)].width=16
ws.cell(row=1,column=1,value='Cross-sectional Conflicted CIO coefficients under alternative performance measures').font=FB
nv=("Each block re-estimates one cross-sectional split with each performance measure. 'Updated' measures use the updated policy benchmark (reported values plus inferred values graded A or B); 'reported only' drops inferred benchmark values; 'document-matched' uses the total-fund return reported on the same page as the benchmark instead of the PPD return; the last row keeps only plan-years in which the PPD and document returns agree within 0.5 percentage points. Specifications, fixed effects and clustering as in the corresponding table.")
ws.merge_cells('A2:E2'); ws['A2']=nv; ws['A2'].font=F; ws['A2'].alignment=JUST; ws.row_dimensions[2].height=70
r=4
for key in ['underpaid','young','priv','fin','rate','fund']:
    sv,cl,(l1,l0),tab=SPL[key]
    ws.cell(row=r,column=1,value=f'{tab}: {l1} vs. {l0}').font=FB; r+=1
    for j,h in enumerate([f'{l1} x Conflicted',f'{l0} x Conflicted','p-value: equal','Observations']):
        c=ws.cell(row=r,column=2+j,value=h); c.font=F; c.alignment=CENTER; c.border=Border(top=thin,bottom=thin)
    ws.cell(row=r,column=1).border=Border(top=thin,bottom=thin); ws.row_dimensions[r].height=30; r+=1
    for yv,yl in YS:
        x=RES[key][yv]; ws.cell(row=r,column=1,value=yl).font=FB if yv in ('peer_adj_ret100','b_AB') else F
        for k,side in enumerate(['a','b']):
            ref=reg(key+' (versions)',yv,side,x[side],x['G']); a=ws.cell(row=r,column=2+k,value=cf(ref)); a.font=F; a.alignment=RIGHT; b=ws.cell(row=r+1,column=2+k,value=tf(ref)); b.font=F; b.alignment=RIGHT
        ref=reg(key+' (versions)',yv,'difference',x['d'],x['G']); c=ws.cell(row=r,column=4,value=pf_(ref)); c.font=F; c.alignment=RIGHT; c.number_format='0.000'
        c=ws.cell(row=r,column=5,value=x['N']); c.font=F; c.alignment=RIGHT; c.number_format='#,##0'
        r+=2
    for c in range(1,6): ws.cell(row=r-1,column=c).border=Border(bottom=thin)
    r+=1
ws.cell(row=r,column=1,value="Note on samples and inference: the young and funding indicators are missing when CIO age or the funding ratio is missing, and the funding indicator no longer requires a non-missing return assumption, so these two splits differ slightly from the published tables; all other peer-adjusted columns reproduce the paper's code exactly. These two-way clustered t-statistics and p-values differ from reghdfe's published ones, in both directions depending on the split; use the attached Stata code for the exact figures.").font=FI
ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=5); ws.cell(row=r,column=1).alignment=JUST; ws.row_dimensions[r].height=45
wb.move_sheet('Estimates', offset=4)
out='/home/user/PensionData/referee_response/time_to_departure/xsec_updated_benchmark_tables_v2.xlsx'; wb.save(out); print('saved')
