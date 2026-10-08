import warnings; warnings.filterwarnings('ignore')
import pandas as pd, numpy as np, pyfixest as pf
from scipy import stats
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

df = pd.read_pickle(__import__('os').path.join(__import__('os').path.dirname(__import__('os').path.abspath(__file__)),'df_do.pkl')).sort_values(['spell','fy']).copy()
df['fee_ratio_pct'] = df.total_fee_pct*100
S = df[(df.cio_dum==1)&(df.fy<=2020)&(df.curr1==0)].copy()        # Table III sample
S['last_fy']=S.groupby('spell').fy.transform('max'); S['max_ten']=S.groupby('spell').cio_tenure.transform('max')
S['ydep']=S.year_departs; S.loc[(S.curr==0)&S.ydep.isna(),'ydep']=S.last_fy
S['ten_dep']=S.max_ten+(S.ydep-S.last_fy).clip(lower=0)
half_ok = (S.curr==0)&S.ten_dep.notna()&(S.ten_dep>0)&S.cio_tenure.notna()
S['second'] = np.where(half_ok, (S.cio_tenure>S.ten_dep/2).astype(float), np.nan)
S['conf_first']  = ((S.conf_pen_vendor==1)&(S.second==0)).astype(int)
S['conf_second'] = ((S.conf_pen_vendor==1)&(S.second==1)).astype(int)
S['conf_undef']  = ((S.conf_pen_vendor==1)&S.second.isna()).astype(int)   # conflicted but no departure/tenure info; kept so the reference group is exactly non-conflicted
print("Table III sample:", len(S), "| conflicted rows:", int(S.conf_pen_vendor.sum()), "= first", int(S.conf_first.sum()), "+ second", int(S.conf_second.sum()), "+ undefined", int(S.conf_undef.sum()))
C1=['l_invret','size','l_fund']; C2=['l_eq','l_fi','l_re','l_pe','l_hf','l_alt','l_comd','l_cash']; C3=['fincenter1']; C4=['l_log_cio_age','l_log_cio_tenure','prior_private_exp','Elite_inst']; C5=['l_sib1','wstate_political']
LAB = {'conf_first':'Conflicted CIO x First Half of Tenure (0/1)','conf_second':'Conflicted CIO x Second Half of Tenure (0/1)','l_invret':'Investment Return (t - 1)','size':'Log (Plan Size) (t-1)','l_fund':'Funding Ratio (t-1)',
 'l_eq':'Equity Allocation (%)','l_fi':'Fixed Income Allocation (%)','l_re':'Real Estate Allocation (%)','l_pe':'Private Equity Allocation (%)','l_hf':'Hedge Fund Allocation (%)','l_alt':'Misc. Alternatives Allocation (%)',
 'l_comd':'Commodities Allocation (%)','l_cash':'Cash Allocation (%)','fincenter1':'Financial Center (0/1)','l_log_cio_age':'Log (CIO Age) (t-1)','l_log_cio_tenure':'Log (CIO Tenure) (t-1)',
 'prior_private_exp':'CIO Prior Private Experience (0/1)','Elite_inst':'CIO Attended Elite College (0/1)','l_sib1':'Separate Investment Board (0/1)','wstate_political':'% Political Appointees on Board'}
VARS = ['conf_first','conf_second']+C1+C2+C3+C4+C5
COLS = [dict(ctrl=C1+C2+C3, fe='fy', plan='No', year='Yes'), dict(ctrl=C1+C2+C3+C4+C5, fe='fy', plan='No', year='Yes'),
        dict(ctrl=C1+C2+C3+C4+C5, fe='ppd_id', plan='Yes', year='No'), dict(ctrl=C1+C2+C3+C4+C5, fe='ppd_id + fy', plan='Yes', year='Yes')]
def psd(m):
    V=np.asarray(m._vcov); w,Q=np.linalg.eigh((V+V.T)/2); return (Q*np.clip(w,0,None))@Q.T
def run(y, col):
    x = ['conf_first','conf_second'] + (['conf_undef'] if S.conf_undef.sum()>0 else [])
    m = pf.feols(f'{y} ~ {" + ".join(x+col["ctrl"])} | {col["fe"]}', data=S, vcov={'CRV1':'ppd_id + fy'})
    V = psd(m); names=list(m._coefnames); b=np.asarray(m._beta_hat)
    est = {n:(float(b[i]), float(np.sqrt(max(V[i,i],0)))) for i,n in enumerate(names)}
    c=np.zeros(len(names)); c[names.index('conf_first')]=1; c[names.index('conf_second')]=-1
    diff=(float(c@b), float(np.sqrt(max(c@V@c,0))))
    return dict(est=est, diff=diff, N=int(m._N), Gp=int(m._G[0]), Gy=int(m._G[1]), adjr2=float(m._adj_r2), col=col)
TABLES = {'S1A':[run('peer_adj_ret100',c) for c in COLS], 'S1B':[run('bench_adj_ret_pct',c) for c in COLS], 'S2A':[run('fee_ratio_pct',c) for c in COLS]}
for tab,regs in TABLES.items():
    print(tab, ' | '.join(f"({i+1}) 1st {r['est']['conf_first'][0]:.3f}({r['est']['conf_first'][0]/r['est']['conf_first'][1]:.2f}) 2nd {r['est']['conf_second'][0]:.3f}({r['est']['conf_second'][0]/r['est']['conf_second'][1]:.2f}) diff p={2*stats.t.sf(abs(r['diff'][0]/r['diff'][1]),min(r['Gp'],r['Gy'])-1):.3f} N={r['N']}" for i,r in enumerate(regs)))

path='/home/user/PensionData/referee_response/time_to_departure/conflicted_halves_tables.xlsx'
wb = load_workbook(path)
for nm in ['Table XI-split','Table XII-split','Estimates-split']:
    if nm in wb.sheetnames: del wb[nm]
TNR='Times New Roman'; SZ=10.5
F=Font(name=TNR,size=SZ); FB=Font(name=TNR,size=SZ,bold=True); FI=Font(name=TNR,size=SZ,italic=True); BLUE=Font(name=TNR,size=SZ,color='0000FF')
thin=Side(style='thin'); CENTER=Alignment(horizontal='center',vertical='center',wrap_text=True); RIGHT=Alignment(horizontal='right'); JUST=Alignment(horizontal='justify',wrap_text=True,vertical='top')
E = wb.create_sheet('Estimates-split')
E['A1']='Estimates behind the split tables (blue = regression output; t and p are formulas). Full Table III sample; SE double clustered by plan and year (PSD-corrected); p-values use the t distribution with min(plan clusters, year clusters) - 1 df.'; E['A1'].font=FI
E.column_dimensions['A'].width=40
META=['Observations','Plan clusters','Year clusters','Adjusted R-squared']
rv={v:4+i for i,v in enumerate(VARS)}; rdiff=4+len(VARS); rm={m:rdiff+2+i for i,m in enumerate(META)}
E['A3']='Variable'; E['A3'].font=FB
for v in VARS: E.cell(row=rv[v],column=1,value=LAB[v]).font=F
E.cell(row=rdiff,column=1,value='First Half minus Second Half').font=F
for m in META: E.cell(row=rm[m],column=1,value=m).font=F
REF={}; k=0
for tab,regs in TABLES.items():
    for i,reg in enumerate(regs):
        c0=2+4*k; k+=1
        for j,h in enumerate(['coef','se','t','p']):
            cc=E.cell(row=3,column=c0+j,value=f'{tab}({i+1}) {h}'); cc.font=FB; cc.alignment=CENTER; E.column_dimensions[L(c0+j)].width=9
        dfree=f'MIN({L(c0)}{rm["Plan clusters"]},{L(c0)}{rm["Year clusters"]})-1'; refs={}
        rows = [(v, reg['est'][v]) for v in VARS if v in reg['est']] + [('__diff__', reg['diff'])]
        for v,(b,se) in rows:
            r = rdiff if v=='__diff__' else rv[v]
            E.cell(row=r,column=c0,value=b).font=BLUE; E.cell(row=r,column=c0+1,value=se).font=BLUE
            E.cell(row=r,column=c0+2,value=f'=IF({L(c0+1)}{r}>0,{L(c0)}{r}/{L(c0+1)}{r},"")').font=F
            E.cell(row=r,column=c0+3,value=f'=IF({L(c0+1)}{r}>0,TDIST(ABS({L(c0+2)}{r}),{dfree},2),"")').font=F
            for j in range(4): E.cell(row=r,column=c0+j).number_format='0.0000'
            refs[v]=dict(coef=f"'Estimates-split'!{L(c0)}{r}", t=f"'Estimates-split'!{L(c0+2)}{r}", p=f"'Estimates-split'!{L(c0+3)}{r}")
        for m,val in [('Observations',reg['N']),('Plan clusters',reg['Gp']),('Year clusters',reg['Gy']),('Adjusted R-squared',reg['adjr2'])]:
            E.cell(row=rm[m],column=c0,value=val).font=BLUE; refs[m]=f"'Estimates-split'!{L(c0)}{rm[m]}"
        REF[(tab,i)]=refs
E.freeze_panes='B4'
def coef_f(ref): return f'=TEXT({ref["coef"]},"0.000")&IF({ref["p"]}="","",IF({ref["p"]}<0.01,"***",IF({ref["p"]}<0.05,"**",IF({ref["p"]}<0.1,"*",""))))'
def t_f(ref): return f'=IF({ref["t"]}="","(n/a)","("&TEXT({ref["t"]},"0.00")&")")'
def panel(ws, r, ptitle, deplab, tab, regs):
    n=len(regs); last=1+n
    if ptitle: ws.cell(row=r,column=1,value=ptitle).font=FB; r+=1
    for j in range(n): cc=ws.cell(row=r,column=2+j,value=f'({j+1})'); cc.font=F; cc.alignment=CENTER; cc.border=Border(top=thin)
    ws.cell(row=r,column=1).border=Border(top=thin); r+=1
    for j in range(n): cc=ws.cell(row=r,column=2+j,value=deplab); cc.font=F; cc.alignment=CENTER; cc.border=Border(bottom=thin)
    ws.cell(row=r,column=1).border=Border(bottom=thin); r+=1
    for v in VARS:
        if not any(v in REF[(tab,j)] for j in range(n)): continue
        ws.cell(row=r,column=1,value=LAB[v]).font=F
        for j in range(n):
            ref=REF[(tab,j)]
            if v in ref:
                cc=ws.cell(row=r,column=2+j,value=coef_f(ref[v])); cc.font=F; cc.alignment=RIGHT
                cc=ws.cell(row=r+1,column=2+j,value=t_f(ref[v])); cc.font=F; cc.alignment=RIGHT
        r+=2
    r+=1
    ws.cell(row=r,column=1,value='Observations').font=F
    for j in range(n): cc=ws.cell(row=r,column=2+j,value=f'={REF[(tab,j)]["Observations"]}'); cc.font=F; cc.alignment=RIGHT; cc.number_format='#,##0'
    r+=1; ws.cell(row=r,column=1,value='Adjusted R-squared').font=F
    for j in range(n): cc=ws.cell(row=r,column=2+j,value=f'={REF[(tab,j)]["Adjusted R-squared"]}'); cc.font=F; cc.alignment=RIGHT; cc.number_format='0.000'
    for lab,key in [('Plan FE','plan'),('Year FE','year')]:
        r+=1; ws.cell(row=r,column=1,value=lab).font=F
        for j,reg in enumerate(regs): cc=ws.cell(row=r,column=2+j,value=reg['col'][key]); cc.font=F; cc.alignment=CENTER
    r+=1; ws.cell(row=r,column=1,value='p-value: First Half = Second Half').font=F
    for j in range(n): cc=ws.cell(row=r,column=2+j,value=f'={REF[(tab,j)]["__diff__"]["p"]}'); cc.font=F; cc.alignment=RIGHT; cc.number_format='0.000'
    for c in range(1,last+1): ws.cell(row=r,column=c).border=Border(bottom=thin)
    return r+2
def table_sheet(name, title, notes, panels):
    ws=wb.create_sheet(name); ws.sheet_view.showGridLines=False; ws.column_dimensions['A'].width=40
    for c in range(2,6): ws.column_dimensions[L(c)].width=16
    ws.cell(row=1,column=1,value=title).font=FB
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=5); ws.cell(row=2,column=1,value=notes).font=F; ws.cell(row=2,column=1).alignment=JUST
    ws.row_dimensions[2].height=15*(len(notes)//95+2); r=4
    for ptitle,deplab,tab in panels: r=panel(ws,r,ptitle,deplab,tab,TABLES[tab])
n_undef = int(S.conf_undef.sum())
undef_txt = (f" {n_undef} conflicted plan-years for which tenure at departure cannot be determined are absorbed by a separate indicator that is not reported." if n_undef>0 else "")
notes1=("This table repeats the regressions of Table III, splitting the Conflicted CIO indicator into two indicators: Conflicted CIO x First Half of Tenure, equal to 1 in the fiscal years in which a conflicted CIO's "
 "tenure is at most one half of the CIO's tenure at departure, and Conflicted CIO x Second Half of Tenure, equal to 1 in the remaining fiscal years of a conflicted CIO's tenure. The omitted category is all "
 "plan-years run by non-conflicted CIOs, so each coefficient is the performance differential of conflicted CIOs in that part of their tenure relative to non-conflicted CIOs. Tenure at departure is the CIO's "
 "tenure in the last fiscal year observed in the plan plus any gap to the departure year." + undef_txt + " The performance measure in Panel A is Peer-Adjusted Return, which is the difference of a plan's "
 "return and the average pension plan return each year. The performance measure in Panel B is Benchmark-Adjusted Return, which is calculated as the plan's return minus the return of its policy benchmark, "
 "when available. The other independent variables are as defined in Table I. The last row reports the p-value of a test that the first-half and second-half coefficients are equal. The standard errors are "
 "double clustered by plan and year. Coefficients marked with ***, **, and * are significant at the 1%, 5%, and 10% level, respectively.")
table_sheet('Table XI-split','Table XI: Conflicted CIOs: Plan Performance in the First versus Second Half of Tenure',notes1,[('Panel A. Peer-Adjusted Returns','Peer-adj. Ret. (t)','S1A'),('Panel B. Benchmark-Adjusted Returns','Bench-adj. Ret. (t)','S1B')])
notes2=("This table repeats the regressions of Table XI with Fee Ratio as the dependent variable. Fee Ratio is the plan's total investment fees as a percentage of plan assets. Conflicted CIO x First Half of Tenure "
 "and Conflicted CIO x Second Half of Tenure are defined as in Table XI, and the omitted category is all plan-years run by non-conflicted CIOs. The other independent variables are as defined in Table I. The last "
 "row reports the p-value of a test that the first-half and second-half coefficients are equal. The standard errors are double clustered by plan and year. Coefficients marked with ***, **, and * are significant "
 "at the 1%, 5%, and 10% level, respectively.")
table_sheet('Table XII-split','Table XII: Conflicted CIOs: Fee Ratio in the First versus Second Half of Tenure',notes2,[('','Fee Ratio (t)','S2A')])
# order: split tables first
order = ['Table XI-split','Table XII-split','Table XI','Table XII','Table XIII','Estimates-split','Estimates']
wb._sheets = [wb[n] for n in order if n in wb.sheetnames] + [s for s in wb._sheets if s.title not in order]
wb.save(path); print('saved')
