import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
R=pd.read_pickle('xsec_narrow_res.pkl'); RES=R['RES']; SPL={k:(sv,cl,labs,tab) for k,sv,cl,labs,tab in R['SPL']}; VERS=R['VERS']; YS=R['YS']
C=pd.read_pickle('comp_splits_res.pkl'); CRES=C['CRES']; CS=C['CS']; CVERS=C['CVERS']; SH=C['SH']; SHP=C['SHP']
PR=pd.read_pickle('period_split_res.pkl')
TNR='Times New Roman'; SZ=10.5
F=Font(name=TNR,size=SZ); FB=Font(name=TNR,size=SZ,bold=True); FI=Font(name=TNR,size=SZ,italic=True); BLUE=Font(name=TNR,size=SZ,color='0000FF')
thin=Side(style='thin'); CENTER=Alignment(horizontal='center',vertical='center',wrap_text=True); RIGHT=Alignment(horizontal='right'); LEFT=Alignment(horizontal='left',vertical='center',wrap_text=True); JUST=Alignment(horizontal='justify',wrap_text=True,vertical='top')
YL={'peer_adj_ret100':'Peer-Adjusted Return (t)','b_AB':'Benchmark-Adjusted Return, Updated Benchmark (t)'}
wb=Workbook(); E=wb.active; E.title='Estimates'
E['A1']='Estimates behind the tables. Blue = regression output (coefficient, SE, minimum cluster count). t and p are formulas; p uses the t distribution with (min clusters - 1) df. SE double clustered by plan and fiscal year (plan and fiscal-year-end date for the financial-center split), with the eigenvalue fix.'; E['A1'].font=FI
for j,h in enumerate(['Table','Specification','Dependent variable','Term','Coefficient','SE','Clusters (min)','t','p'],1): E.cell(row=3,column=j,value=h).font=FB
for c,w in zip('ABCD',[30,48,22,40]): E.column_dimensions[c].width=w
er=[4]
def reg(tab,spec,y,term,r,G):
    i=er[0]; er[0]+=1
    for j,v in enumerate([tab,spec,y,term],1): E.cell(row=i,column=j,value=v).font=F
    E.cell(row=i,column=5,value=r['est']).font=BLUE; E.cell(row=i,column=6,value=r['se']).font=BLUE; E.cell(row=i,column=7,value=G).font=BLUE
    E.cell(row=i,column=8,value=f'=IF(F{i}>0,E{i}/F{i},"")').font=F; E.cell(row=i,column=9,value=f'=IF(H{i}="","",TDIST(ABS(H{i}),G{i}-1,2))').font=F
    for j in [5,6,8,9]: E.cell(row=i,column=j).number_format='0.0000'
    return dict(coef=f'Estimates!E{i}',t=f'Estimates!H{i}',p=f'Estimates!I{i}')
def cf(ref): return f'=IF({ref["p"]}="","n/a",TEXT({ref["coef"]},"0.000")&IF({ref["p"]}<0.01,"***",IF({ref["p"]}<0.05,"**",IF({ref["p"]}<0.1,"*",""))))'
def tf(ref): return f'=IF({ref["t"]}="","","("&TEXT({ref["t"]},"0.00")&")")'
def pf_(ref): return f'=IF({ref["p"]}="","",{ref["p"]})'
def put(ws,r,c,ref,fnt=F):
    a=ws.cell(row=r,column=c,value=cf(ref)); a.font=fnt; a.alignment=RIGHT; b=ws.cell(row=r+1,column=c,value=tf(ref)); b.font=fnt; b.alignment=RIGHT
def setup(name,title,txt,ncol,widths,per=105):
    ws=wb.create_sheet(name); ws.sheet_view.showGridLines=False
    for j,w in enumerate(widths): ws.column_dimensions[L(1+j)].width=w
    ws.cell(row=1,column=1,value=title).font=FB
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=ncol); ws['A2']=txt; ws['A2'].font=F; ws['A2'].alignment=JUST; ws.row_dimensions[2].height=14*(len(txt)//per+2)
    ws.page_setup.orientation='landscape'; ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    return ws
NARROWTXT=("The control variables are Size, the lagged funding ratio and the lagged allocations to equities, fixed income, real estate, private equity, hedge funds, other alternatives, commodities and cash, each interacted with the group indicator (the funding split omits the lagged funding ratio, as in the paper's code). ")
BENCH=("Benchmark-Adjusted Return is the plan's return minus the updated policy benchmark (reported in the plan's financial report or, where none is stated, an inferred value graded A or B). ")
TAIL="The p-value row tests that the two Conflicted CIO coefficients are equal. The standard errors are double clustered by plan and year (plan and fiscal-year-end date for the financial-center split). Coefficients marked with ***, **, and * are significant at the 1%, 5%, and 10% level, respectively."
MAINV='narrow_fye'
PAPER_FIRST_G0={'priv','fin'}
def panel(ws,r,key,ptitle):
    sv,cl,(l1,l0),tab=SPL[key]; order=[('b',l0),('a',l1)] if key in PAPER_FIRST_G0 else [('a',l1),('b',l0)]
    ws.cell(row=r,column=1,value=ptitle).font=FB; r+=1
    for j,(y,_) in enumerate(YS):
        c0=2+2*j; ws.merge_cells(start_row=r,start_column=c0,end_row=r,end_column=c0+1); c=ws.cell(row=r,column=c0,value=YL[y]); c.font=F; c.alignment=CENTER
        for cc in (c0,c0+1): ws.cell(row=r,column=cc).border=Border(top=thin)
    ws.cell(row=r,column=1).border=Border(top=thin); ws.row_dimensions[r].height=30; r+=1
    for j,_ in enumerate(YS):
        for k,(side,lab) in enumerate(order): c=ws.cell(row=r,column=2+2*j+k,value=lab); c.font=F; c.alignment=CENTER; c.border=Border(bottom=thin)
    ws.cell(row=r,column=1).border=Border(bottom=thin); ws.row_dimensions[r].height=30; r+=1
    ws.cell(row=r,column=1,value='Conflicted CIO').font=F
    for j,(y,_) in enumerate(YS):
        x=RES[MAINV][key][y]
        for k,(side,lab) in enumerate(order): put(ws,r,2+2*j+k,reg(tab,'Narrow controls, fye1 FE',y,f'{lab} x Conflicted',x[side],x['G']))
    r+=3
    rows=[('p-value: equal coefficients',None,'0.000'),('Observations',lambda x:x['N'],'#,##0'),('Adjusted R-squared',lambda x:x['adjr2'],'0.000')]
    for lab,fn,fmt in rows:
        ws.cell(row=r,column=1,value=lab).font=F
        for j,(y,_) in enumerate(YS):
            x=RES[MAINV][key][y]; c0=2+2*j; ws.merge_cells(start_row=r,start_column=c0,end_row=r,end_column=c0+1)
            v=pf_(reg(tab,'Narrow controls, fye1 FE',y,'difference',x['d'],x['G'])) if fn is None else fn(x)
            c=ws.cell(row=r,column=c0,value=v); c.font=F; c.alignment=CENTER; c.number_format=fmt
        r+=1
    for lab,val in [('Controls','Narrow'),('Fiscal-Year-End FE','Yes')]:
        ws.cell(row=r,column=1,value=lab).font=F
        for j in range(len(YS)):
            c0=2+2*j; ws.merge_cells(start_row=r,start_column=c0,end_row=r,end_column=c0+1); c=ws.cell(row=r,column=c0,value=val); c.font=F; c.alignment=CENTER
        r+=1
    for c in range(1,6): ws.cell(row=r-1,column=c).border=Border(bottom=thin)
    return r+2
COMMON=("Each panel reports one regression of plan performance on Conflicted CIO interacted with the group indicator, with the control variables also interacted with the group indicator, and fiscal-year-end date fixed effects. "+NARROWTXT+BENCH+TAIL)
for name,title,intro,pans in [
  ('Table V','Table V: CIO Incentives and Plan Performance, Narrow Controls',"Underpaid is the paper's indicator (CIO total compensation below $224,578, the pooled 2009-2014 median; see xsec_interactions_fye1.xlsx, Underpaid checks). Panel C instead uses pay below the median of reported CIO pay in the same fiscal year, filled within the CIO's spell at the plan. Young CIOs are 46 years old or younger (missing when CIO age is missing). ",
   [('underpaid','Panel A. Effect of Conflict for Underpaid CIOs'),('young','Panel B. Effect of Conflict for Younger CIOs'),('up_ann','Panel C. Effect of Conflict for Underpaid CIOs, Annual Median')]),
  ('Table VI','Table VI: Alternative Pathways and Plan Performance, Narrow Controls',"Prior Private Exp. equals 1 for CIOs with prior private firm experience; Financial Center equals 1 for plans headquartered in financial centers. ",
   [('priv','Panel A. Effect of Conflict for CIOs With and Without Prior Private Experience'),('fin','Panel B. Effect of Conflict for Plans Located in Financial Centers')]),
  ('Table VII','Table VII: Governance and Plan Performance, Narrow Controls',"High Discount Rate equals 1 if the plan's discount rate is above 8%; High Funding equals 1 for funding ratios of at least 92% (missing when the funding ratio is missing). ",
   [('rate','Panel A. Effect of Conflict Based on Liability Discount Rate'),('fund','Panel B. Effect of Conflict Based on Funding')])]:
    ws=setup(name,title,intro+COMMON,5,[30,15,15,15,15])
    r=4
    for key,pt in pans: r=panel(ws,r,key,pt)
# ---------- specification comparison ----------
VL={a:b for a,b,_ in VERS}
ws=setup('Specification comparison','Cross-Sectional Splits Under Alternative Controls and Fixed Effects',
 "Each block re-estimates one split under five specifications. 'Narrow controls' are listed in the table notes of Tables V-VII on this workbook; 'Paper controls' are the control variables of the corresponding regression in the paper's code. 'Paper-controls sample' keeps only plan-years with all of the paper's controls non-missing, so the row differs from the one above it only in the control variables. The last row of each block is the published specification. "+BENCH+TAIL,10,[52,13,13,9,8,3,13,13,9,8],110)
r=4
hdr=['','Group 1 x Conflicted','Group 0 x Conflicted','p-value equal','Obs.','','Group 1 x Conflicted','Group 0 x Conflicted','p-value equal','Obs.']
for key in ['underpaid','up_ann','young','priv','fin','rate','fund']:
    sv,cl,(l1,l0),tab=SPL[key]
    ws.cell(row=r,column=1,value=f'{tab}: Group 1 = {l1}; Group 0 = {l0}').font=FB; r+=1
    for j,(y,_) in enumerate(YS):
        c0=2+5*j; ws.merge_cells(start_row=r,start_column=c0,end_row=r,end_column=c0+3); c=ws.cell(row=r,column=c0,value=YL[y]); c.font=F; c.alignment=CENTER
    r+=1
    for j,h in enumerate(hdr):
        if h: c=ws.cell(row=r,column=1+j,value=h); c.font=F; c.alignment=CENTER
        ws.cell(row=r,column=1+j).border=Border(top=thin,bottom=thin)
    ws.row_dimensions[r].height=30; r+=1
    for v,vl,_ in VERS:
        ws.cell(row=r,column=1,value=vl).font=FB if v==MAINV else F; ws.cell(row=r,column=1).alignment=LEFT
        for j,(y,_) in enumerate(YS):
            x=RES[v][key][y]; c0=2+5*j
            for k,side in enumerate(['a','b']): put(ws,r,c0+k,reg(tab,vl,y,side,x[side],x['G']))
            c=ws.cell(row=r,column=c0+2,value=pf_(reg(tab,vl,y,'difference',x['d'],x['G']))); c.font=F; c.alignment=RIGHT; c.number_format='0.000'
            c=ws.cell(row=r,column=c0+3,value=x['N']); c.font=F; c.alignment=RIGHT; c.number_format='#,##0'
        r+=2
    for c in range(1,11): ws.cell(row=r-1,column=c).border=Border(bottom=thin)
    r+=1
# ---------- compensation splits ----------
ws=setup('Compensation splits','Effect of Conflict by CIO Pay: Alternative Definitions',
 "Each row splits the sample by a different definition of high versus low CIO pay, computed from reported CIO_TotalComp over all plan-years with reported pay and filled within the CIO's spell at the plan (years without reported pay take the nearest earlier reported year, or the first reported year), as the paper's Underpaid measure is filled. "
 "The first row reproduces the paper's Underpaid cutoff on this sample (spells without any reported pay drop out). Inflation-adjusted pay deflates by the CPI-U annual average (BLS series CUUR0000SA0) to 2020 dollars. Pay given plan size is the residual from a regression of log pay on Size and fiscal-year indicators. The top-third, top-quarter and bottom-third cutoffs are computed within each fiscal year. "
 "The last columns report the share of plan-years in the lower-paid group, overall, for conflicted plan-years, and by period (2001-2008 / 2009-2014 / 2015-2020). Panel A uses the narrow controls with fiscal-year-end date fixed effects; Panel B uses the paper's controls for the Underpaid split with year fixed effects (the published specification). "+BENCH+TAIL,13,[46,13,13,9,8,3,13,13,9,8,3,10,18],120)
r=4
for v,vl,_ in CVERS:
    ws.cell(row=r,column=1,value=f"Panel {'A' if v=='narrow_fye' else 'B'}. {vl}").font=FB; r+=1
    for j,(y,_) in enumerate(YS):
        c0=2+5*j; ws.merge_cells(start_row=r,start_column=c0,end_row=r,end_column=c0+3); c=ws.cell(row=r,column=c0,value=YL[y]); c.font=F; c.alignment=CENTER
    r+=1
    h2=['Definition of the lower-paid group','Lower-paid x Conflicted','Higher-paid x Conflicted','p-value equal','Obs.','','Lower-paid x Conflicted','Higher-paid x Conflicted','p-value equal','Obs.','','Share lower-paid (all / conflicted)','Share lower-paid by period']
    for j,h in enumerate(h2):
        if h: c=ws.cell(row=r,column=1+j,value=h); c.font=F; c.alignment=CENTER
        ws.cell(row=r,column=1+j).border=Border(top=thin,bottom=thin)
    ws.row_dimensions[r].height=45; r+=1
    for k,lab,l1,l0 in CS:
        ws.cell(row=r,column=1,value=lab).font=F; ws.cell(row=r,column=1).alignment=LEFT
        for j,(y,_) in enumerate(YS):
            x=CRES[v][k][y]; c0=2+5*j
            for kk,side in enumerate(['a','b']): put(ws,r,c0+kk,reg('Compensation: '+lab,vl,y,side,x[side],x['G']))
            c=ws.cell(row=r,column=c0+2,value=pf_(reg('Compensation: '+lab,vl,y,'difference',x['d'],x['G']))); c.font=F; c.alignment=RIGHT; c.number_format='0.000'
            c=ws.cell(row=r,column=c0+3,value=x['N']); c.font=F; c.alignment=RIGHT; c.number_format='#,##0'
        c=ws.cell(row=r,column=12,value=f"{SH[k][0]:.2f} / {SH[k][1]:.2f}"); c.font=F; c.alignment=RIGHT
        c=ws.cell(row=r,column=13,value=' / '.join(f'{z:.2f}' for z in SHP[k])); c.font=F; c.alignment=RIGHT
        r+=2
    for c in range(1,14): ws.cell(row=r-1,column=c).border=Border(bottom=thin)
    r+=2
# ---------- period split ----------
ws=setup('Period split','Effect of Conflict in 2001-2008 and 2009-2020',
 "Each regression interacts Conflicted CIO and the control variables with an indicator for fiscal years 2001-2008, as in the cross-sectional splits. The first specification uses the narrow controls with fiscal-year-end date fixed effects; the second uses the paper's controls for the Underpaid split with year fixed effects. The sample has 192 conflicted plan-years (32 CIOs) in 2001-2008 and 175 (27 CIOs) in 2009-2020. "+BENCH+TAIL,10,[46,13,13,9,8,3,13,13,9,8],110)
r=4
for j,(y,_) in enumerate(YS):
    c0=2+5*j; ws.merge_cells(start_row=r,start_column=c0,end_row=r,end_column=c0+3); c=ws.cell(row=r,column=c0,value=YL[y]); c.font=F; c.alignment=CENTER
r+=1
for j,h in enumerate(['Specification','2001-2008 x Conflicted','2009-2020 x Conflicted','p-value equal','Obs.','','2001-2008 x Conflicted','2009-2020 x Conflicted','p-value equal','Obs.']):
    if h: c=ws.cell(row=r,column=1+j,value=h); c.font=F; c.alignment=CENTER
    ws.cell(row=r,column=1+j).border=Border(top=thin,bottom=thin)
ws.row_dimensions[r].height=30; r+=1
for v,vl in [('narrow_fye','Narrow controls, fiscal-year-end FE'),('paper_fy','Paper controls, year FE (published specification)')]:
    ws.cell(row=r,column=1,value=vl).font=F; ws.cell(row=r,column=1).alignment=LEFT
    for j,(y,_) in enumerate(YS):
        x=PR[v][y]; c0=2+5*j
        for kk,side in enumerate(['a','b']): put(ws,r,c0+kk,reg('Period split',vl,y,side,x[side],x['G']))
        c=ws.cell(row=r,column=c0+2,value=pf_(reg('Period split',vl,y,'difference',x['d'],x['G']))); c.font=F; c.alignment=RIGHT; c.number_format='0.000'
        c=ws.cell(row=r,column=c0+3,value=x['N']); c.font=F; c.alignment=RIGHT; c.number_format='#,##0'
    r+=2
for c in range(1,11): ws.cell(row=r-1,column=c).border=Border(bottom=thin)

# ---------- log pay x time ----------
CT=pd.read_pickle('comp_time_check.pkl'); CTR=CT['CT']; CTV=CT['V']
ws=setup('Log pay x time','Log CIO Pay Interaction, Allowing the Effect of Conflict to Vary Over Time',
 "Each column re-estimates the Log CIO Total Comp. column of the interaction specification (xsec_interactions_fye1.xlsx: all Table III controls, fiscal-year-end date fixed effects, log pay in the same fiscal year centered at its regression-sample mean, plan-years with reported pay). "
 "Columns (2)-(4) add interactions of Conflicted CIO with an indicator for 2001-2008, a linear fiscal-year trend, or fiscal-year indicators (each centered), so that the effect of conflict may differ over time; in column (4) the log pay interaction is identified only from differences in pay among conflicted CIOs in the same fiscal year. "+BENCH+"The standard errors are double clustered by plan and year. Coefficients marked with ***, **, and * are significant at the 1%, 5%, and 10% level, respectively.",5,[40,18,18,18,18],100)
r=4
for y,ptitle in [('peer_adj_ret100','Panel A. Peer-Adjusted Returns'),('b_AB','Panel B. Benchmark-Adjusted Returns (updated benchmark)')]:
    ws.cell(row=r,column=1,value=ptitle).font=FB; r+=1
    for j,(v,vl) in enumerate(CTV): c=ws.cell(row=r,column=2+j,value=f'({j+1})\n{vl}'); c.font=F; c.alignment=CENTER; c.border=Border(top=thin,bottom=thin)
    ws.cell(row=r,column=1).border=Border(top=thin,bottom=thin); ws.row_dimensions[r].height=58; r+=1
    for term,tl in [('conf','Conflicted CIO (at mean log pay)'),('confX','Conflicted CIO x Log CIO Total Comp.'),('X','Log CIO Total Comp.'),('early','Conflicted CIO x 2001-2008')]:
        ws.cell(row=r,column=1,value=tl).font=F
        for j,(v,vl) in enumerate(CTV):
            x=CTR[v][y]
            if term=='early':
                if 'early' not in x['ext']: continue
                est=x['ext']['early']
            else: est=x[term]
            put(ws,r,2+j,reg('Log pay x time',vl,y,tl,est,x['G']))
        r+=2
    for lab,fn,fmt in [('Observations',lambda x:x['N'],'#,##0'),('Adjusted R-squared',lambda x:x['adjr2'],'0.000')]:
        ws.cell(row=r,column=1,value=lab).font=F
        for j,(v,vl) in enumerate(CTV): c=ws.cell(row=r,column=2+j,value=fn(CTR[v][y])); c.font=F; c.alignment=RIGHT; c.number_format=fmt
        r+=1
    for c in range(1,6): ws.cell(row=r-1,column=c).border=Border(bottom=thin)
    r+=2
wb.move_sheet('Estimates', offset=len(wb.sheetnames)-1)
out='/home/user/PensionData/referee_response/time_to_departure/xsec_narrow_controls.xlsx'; wb.save(out); print('saved',out)
