import pandas as pd, numpy as np
from scipy import stats
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
D=pd.read_pickle('table3_conflicted_halves.pkl'); RES=D['RES']; YS=D['YS']; NC=D['NC']
TNR='Times New Roman'; SZ=10.5
F=Font(name=TNR,size=SZ); FB=Font(name=TNR,size=SZ,bold=True); FI=Font(name=TNR,size=SZ,italic=True)
thin=Side(style='thin'); CENTER=Alignment(horizontal='center',vertical='center',wrap_text=True); RIGHT=Alignment(horizontal='right'); LEFT=Alignment(horizontal='left',vertical='center',wrap_text=True); JUST=Alignment(horizontal='justify',wrap_text=True,vertical='top')
def st(p): return '***' if p<0.01 else '**' if p<0.05 else '*' if p<0.1 else ''
LAB={'h1':'Conflicted CIO x First Half of Tenure (0/1)','h2':'Conflicted CIO x Second Half of Tenure (0/1)','hu':'Conflicted CIO, halves undefined (0/1)',
 'l_invret':'Investment Return (t-1)','size':'Log (Plan Size) (t-1)','l_fund':'Funding Ratio (t-1)','l_eq':'Equity Allocation (t-1)','l_fi':'Fixed Income Allocation (t-1)','l_re':'Real Estate Allocation (t-1)',
 'l_pe':'Private Equity Allocation (t-1)','l_hf':'Hedge Fund Allocation (t-1)','l_alt':'Misc. Alternatives Allocation (t-1)','l_comd':'Commodities Allocation (t-1)','l_cash':'Cash Allocation (t-1)','fincenter1':'Financial Center (0/1)',
 'l_log_cio_age':'Log (CIO Age) (t-1)','l_log_cio_tenure':'Log (CIO Tenure) (t-1)','prior_private_exp':'CIO Prior Private Experience (0/1)','Elite_inst':'CIO Attended Elite College (0/1)','l_sib1':'Separate Investment Board (0/1)','wstate_political':'% Political Appointees on Board'}
c=NC['conf_pen_vendor']
BASE=("This table repeats Table III with Conflicted CIO split into two indicators: Conflicted CIO x First Half of Tenure equals 1 in a conflicted CIO's plan-years in which tenure at the plan is at most half of tenure at departure, and Conflicted CIO x Second Half of Tenure equals 1 in the remaining conflicted plan-years. "
      f"Of the {c['h1']+c['h2']+c['hu']} conflicted plan-years in the Table III sample, {c['h1']} are in the first half and {c['h2']} in the second half; the {c['hu']} conflicted plan-years without usable tenure information (tenure missing or zero) lack the lagged controls and drop out of every regression, so the reference group is exactly the non-conflicted plan-years. "
      "Columns 5 and 6 split the last-5-year and last-10-year Conflicted CIO definitions the same way. The performance measure in Panel A is Peer-Adjusted Return; in Panel B it is Benchmark-Adjusted Return (plan return minus its policy benchmark: reported, else inferred grade A or B). "
      "'Second minus first' is the difference between the two half coefficients, with its p-value. Other variables are as defined in Table I. t-statistics in parentheses; standard errors double clustered by plan and year. ***, **, and * denote significance at the 1%, 5%, and 10% level.")
SHEETS=[('fy','Year FE','Table III with Conflicted CIO Split Into the First and Second Half of Tenure',''),
        ('fye','Fiscal-year-end FE','Table III with Conflicted CIO Split Into the First and Second Half of Tenure: Fiscal-Year-End Date Fixed Effects',' Year fixed effects are replaced by fiscal-year-end date fixed effects.'),
        ('fy_noten','No tenure or age','Table III with Conflicted CIO Split Into the First and Second Half of Tenure: Without Log CIO Tenure and Log CIO Age',' Log (CIO Tenure) and Log (CIO Age) are dropped from the controls because tenure rises mechanically from the first to the second half.')]
wb=Workbook(); first=True
for ver,sname,title,extra in SHEETS:
    ws=wb.active if first else wb.create_sheet(); first=False; ws.title=sname; ws.sheet_view.showGridLines=False
    ws.column_dimensions['A'].width=44
    for j in range(2,8): ws.column_dimensions[L(j)].width=15
    ws['A1']=title; ws['A1'].font=FB
    note=BASE+extra; ws.merge_cells('A2:G2'); ws['A2']=note; ws['A2'].font=F; ws['A2'].alignment=JUST; ws.row_dimensions[2].height=14*(len(note)//115+2)
    r=4
    for pi,(y,yl) in enumerate(YS):
        regs=RES[(ver,y)]
        ws.cell(row=r,column=1,value=f"Panel {'AB'[pi]}. {yl}").font=FB; r+=1
        for j in range(6): x=ws.cell(row=r,column=2+j,value=f'({j+1})'); x.font=F; x.alignment=CENTER; x.border=Border(top=thin)
        ws.cell(row=r,column=1).border=Border(top=thin); r+=1
        for j in range(6): x=ws.cell(row=r,column=2+j,value=('Peer-adj. Ret. (t)' if pi==0 else 'Bench-adj. Ret. (t)')); x.font=F; x.alignment=CENTER; x.border=Border(bottom=thin)
        ws.cell(row=r,column=1).border=Border(bottom=thin); r+=1
        for v,l in LAB.items():
            if not any(v in x['est'] for x in regs): continue
            ws.cell(row=r,column=1,value=l).font=FB if v in ('h1','h2') else F; ws.cell(row=r,column=1).alignment=LEFT
            for j,x in enumerate(regs):
                if v not in x['est']: continue
                b,se=x['est'][v]; t=b/se; p=2*stats.t.sf(abs(t),x['G']-1)
                a=ws.cell(row=r,column=2+j,value=f"{b:.3f}{st(p)}"); a.font=FB if v in ('h1','h2') else F; a.alignment=RIGHT
                a=ws.cell(row=r+1,column=2+j,value=f"({t:.2f})"); a.font=F; a.alignment=RIGHT
            r+=2
        ws.cell(row=r,column=1,value='Second minus first half').font=FI
        ws.cell(row=r+1,column=1,value='  p-value').font=FI
        for j,x in enumerate(regs):
            b,se=x['diff']; t=b/se; p=2*stats.t.sf(abs(t),x['G']-1)
            a=ws.cell(row=r,column=2+j,value=f"{b:.3f}{st(p)}"); a.font=FI; a.alignment=RIGHT
            a=ws.cell(row=r+1,column=2+j,value=f"{p:.3f}"); a.font=FI; a.alignment=RIGHT
        r+=3
        for lab,vals in [('Observations',[f"{x['N']:,}" for x in regs]),('Adjusted R-squared',[f"{x['adjr2']:.3f}" for x in regs]),
                         ('Plan FE',[x['col']['plan'] for x in regs]),('Fiscal-year-end FE' if ver=='fye' else 'Year FE',[x['col']['year'] for x in regs]),
                         ('Conflicted CIO defined over',[('Last '+x['col']['lastn']+' years') if x['col']['lastn'] else 'All years' for x in regs])]:
            ws.cell(row=r,column=1,value=lab).font=F
            for j,v in enumerate(vals): a=ws.cell(row=r,column=2+j,value=v); a.font=F; a.alignment=RIGHT if lab in ('Observations','Adjusted R-squared') else CENTER
            r+=1
        for cc in range(1,8): ws.cell(row=r-1,column=cc).border=Border(bottom=thin)
        r+=2
    ws.page_setup.orientation='portrait'; ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
out='/home/user/PensionData/referee_response/time_to_departure/table3_conflicted_halves.xlsx'; wb.save(out); print('saved',out)
