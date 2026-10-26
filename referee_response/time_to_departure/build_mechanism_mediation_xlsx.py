import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
D=pd.read_pickle('mechanism_mediation.pkl'); MED=D['MED']; CH=D['CH']; LEV=D['LEV']; HALF=D['HALF']; EW=D['EW']; FSL=D['FSL']; YS=D['YS']; OUTS=D['OUTS']; MCOLS=D['MCOLS']; NMS=D['NMS']
SL=pd.read_pickle('returns_spell_slope.pkl')['A']; AF=pd.read_pickle('allocations_fees_halves_equal_weight.pkl')['SL']
TNR='Times New Roman'; SZ=10.5
F=Font(name=TNR,size=SZ); FB=Font(name=TNR,size=SZ,bold=True); FI=Font(name=TNR,size=SZ,italic=True)
thin=Side(style='thin'); CENTER=Alignment(horizontal='center',vertical='center',wrap_text=True); RIGHT=Alignment(horizontal='right',vertical='center'); RIGHTW=Alignment(horizontal='right',vertical='center',wrap_text=True); LEFT=Alignment(horizontal='left',vertical='center',wrap_text=True); JUST=Alignment(horizontal='justify',wrap_text=True,vertical='top')
def st(p): return '' if p is None or np.isnan(p) else '***' if p<0.01 else '**' if p<0.05 else '*' if p<0.1 else ''
def put(ws,r,c,v,font=F,al=RIGHT):
    x=ws.cell(row=r,column=c,value=v); x.font=font; x.alignment=al; return x
def rule(ws,r,n,top=False):
    for c in range(1,n+1): ws.cell(row=r,column=c).border=Border(top=thin) if top else Border(bottom=thin)
wb=Workbook()
# ======================= Sheet 1: mediation =======================
ws=wb.active; ws.title='Mediation'; ws.sheet_view.showGridLines=False
NC=1+len(MCOLS); ws.column_dimensions['A'].width=40
for j in range(2,NC+1): ws.column_dimensions[L(j)].width=15
ws['A1']='Does the Fee or Allocation Channel Absorb the Conflicted CIO Coefficient?'; ws['A1'].font=FB
note=("Each column is a Table III regression of returns on Conflicted CIO with fiscal-year-end date fixed effects (Panels A1 and B1) or plan and fiscal-year-end date fixed effects (Panels A2 and B2), on a common sample of plan-years with non-missing Fee Ratio, Fee Slope and alternatives share, so the Conflicted coefficient is comparable across columns. "
      "Column 1 omits the lagged allocation controls; column 2 has the Table III controls; columns 3-6 add the candidate mediators: Fee Ratio (total external fees / assets, %), Fee Slope (the CIO-plan fee trend used in Table IX), the alternatives share of assets (real estate, private equity, hedge funds, other alternatives and commodities, %), and all three. "
      "'Change in Conflicted vs. column 2' is the column 2 coefficient minus this column's (positive = the mediator absorbs part of the underperformance), with a 95% interval and p-value from a plan-block bootstrap ("+str(D['B'])+" draws). "
      "'Path a' rows regress each mediator on Conflicted CIO with the same controls and fixed effects (alternatives share without the allocation controls). The mediation sample requires fee data, so it is smaller than the Table III sample; on it, the benchmark-adjusted Conflicted coefficient with plan and fiscal-year-end fixed effects is small and insignificant (Panel B2), so that panel has little to mediate. t-statistics in parentheses; standard errors double clustered by plan and year. ***, **, and * denote significance at the 1%, 5%, and 10% level.")
ws.merge_cells(f'A2:{L(NC)}2'); ws['A2']=note; ws['A2'].font=F; ws['A2'].alignment=JUST; ws.row_dimensions[2].height=14*(len(note)//125+1)
r=4
for pi,(y,yl) in enumerate(YS):
    for fi,(fk,flab) in enumerate([('fye','Fiscal-year-end FE'),('plan','Plan + fiscal-year-end FE')]):
        n,nc,ncio=NMS[(y,fk)]
        ws.cell(row=r,column=1,value=f"Panel {'AB'[pi]}{fi+1}. {yl}, {flab}").font=FB; r+=1
        put(ws,r,1,'',F,LEFT)
        for j,(ck,lab) in enumerate(MCOLS): put(ws,r,2+j,f"({j+1})\n{lab}",F,CENTER)
        rule(ws,r,NC,top=True); ws.row_dimensions[r].height=44; rule(ws,r,NC); r+=1
        for v,lab in [('conf_pen_vendor','Conflicted CIO (0/1)'),('fee','Fee Ratio (%)'),('fslope','Fee Slope'),('alt','Alternatives share (%)')]:
            put(ws,r,1,lab,FB if v=='conf_pen_vendor' else F,LEFT)
            for j,(ck,_) in enumerate(MCOLS):
                c=MED[(y,fk,ck)]['coef']
                if v in c:
                    b,se,p=c[v]; put(ws,r,2+j,f"{b:.3f}{st(p)}",FB if v=='conf_pen_vendor' else F); put(ws,r+1,2+j,f"({b/se:.2f})")
            r+=2
        put(ws,r,1,'Change in Conflicted vs. column 2',FI,LEFT); put(ws,r+1,1,'  95% bootstrap interval',F,LEFT); put(ws,r+2,1,'  bootstrap p-value',F,LEFT)
        for j,(ck,_) in enumerate(MCOLS):
            if ck=='t3': continue
            q=CH[(y,fk,ck)]; put(ws,r,2+j,f"{q['diff']:+.3f}",FI); put(ws,r+1,2+j,f"[{q['lo']:+.2f}, {q['hi']:+.2f}]"); put(ws,r+2,2+j,f"{q['p']:.3f}")
        r+=3
        put(ws,r,1,'Observations',F,LEFT)
        for j,(ck,_) in enumerate(MCOLS): put(ws,r,2+j,f"{MED[(y,fk,ck)]['N']:,}")
        r+=1; put(ws,r,1,'Adjusted R-squared',F,LEFT)
        for j,(ck,_) in enumerate(MCOLS): put(ws,r,2+j,f"{MED[(y,fk,ck)]['adjr2']:.3f}")
        r+=1; put(ws,r,1,f'Conflicted plan-years / conflicted CIOs: {nc} / {ncio}',FI,LEFT); r+=1
        for med,lab in [('fee','Fee Ratio'),('fslope','Fee Slope'),('alt','Alternatives share')]:
            b,se,p=MED[(y,fk,'a_'+med)]['coef']['conf_pen_vendor']
            put(ws,r,1,f'Path a: {lab} on Conflicted CIO (t)',FI,LEFT); put(ws,r,2,f"{b:+.3f}{st(p)} ({b/se:.2f})",FI); r+=1
        rule(ws,r-1,NC); r+=1
ws.page_setup.orientation='landscape'; ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
# ======================= Sheet 2: mechanism summary =======================
ws=wb.create_sheet('Mechanism summary'); ws.sheet_view.showGridLines=False
W=[50,21,21,21,21]
for j,w in enumerate(W): ws.column_dimensions[L(1+j)].width=w
ws['A1']='Steering or Flat? Level, Timing and Mediation Evidence in One Place'; ws['A1'].font=FB
note2=("Panel A: Conflicted CIO coefficient (Table III controls; alternatives share without allocation controls) and the CIO-plan Fee Slope comparison. "
       "Panel B: does the effect build up toward departure? Rows 1-2 split Conflicted CIO into the first and second half of the CIO's tenure (second minus first, p-value). Row 3 is the equal-weighted within-CIO change from the first to the second half for departed CIOs, conflicted minus other departing CIOs, with CIO and fiscal-year-end fixed effects (each CIO-plan's halves weighted equally). Row 4 is the equal-weighted difference in spell-level OLS slopes on fiscal year (conflicted minus other departing CIOs; at least three years; winsorized), with the rank-sum p-value for medians. "
       "Panel C: change in the Conflicted coefficient when the mediators are added (Mediation sheet, column 2 minus column 6) and the bootstrap p-value. Entries are estimates with p-values in parentheses; standard errors double clustered by plan and year unless noted. ***, **, and * denote significance at the 1%, 5%, and 10% level.")
ws.merge_cells('A2:E2'); ws['A2']=note2; ws['A2'].font=F; ws['A2'].alignment=JUST; ws.row_dimensions[2].height=14*(len(note2)//120+1)
r=4; H=['']+[lab for _,lab in OUTS]
for j,h in enumerate(H): put(ws,r,1+j,h,FB,CENTER)
rule(ws,r,5,top=True); ws.row_dimensions[r].height=30; rule(ws,r,5); r+=1
def cpp(e,p,dec=3): return f"{e:+.{dec}f}{st(p)} ({p:.3f})"
ws.cell(row=r,column=1,value='Panel A. Level effect of a Conflicted CIO').font=FB; r+=1
for fk,lab in [('fye','  Conflicted CIO, fiscal-year-end FE'),('plan','  Conflicted CIO, plan + fiscal-year-end FE')]:
    put(ws,r,1,lab,F,LEFT)
    for j,(y,_) in enumerate(OUTS): b,se,p=LEV[(y,fk)]['coef']['conf_pen_vendor']; put(ws,r,2+j,cpp(b,p))
    r+=1
put(ws,r,1,'  Fee Slope (CIO-plan level), conflicted minus other: mean (p); medians (rank-sum p)',F,LEFT); put(ws,r,4,f"{FSL['diff']:+.3f}{st(FSL['p'])} ({FSL['p']:.3f}); medians {FSL['medc']:+.3f} vs {FSL['medo']:+.3f} ({FSL['prs']:.3f})",F,RIGHTW); ws.row_dimensions[r].height=44; r+=1
ws.cell(row=r,column=1,value='Panel B. Does it build up toward departure? (second half minus first half)').font=FB; r+=1
for fk,lab in [('fye','  Conflicted x Second minus x First Half, fiscal-year-end FE'),('plan','  Conflicted x Second minus x First Half, plan + FYE FE')]:
    put(ws,r,1,lab,F,LEFT)
    for j,(y,_) in enumerate(OUTS): e,se,p=HALF[(y,fk)]['diff']; put(ws,r,2+j,cpp(e,p))
    r+=1
put(ws,r,1,'  Within-CIO halves change, conflicted minus other (equal weight, CIO + FYE FE)',F,LEFT)
for j,(y,_) in enumerate(OUTS): q=EW[y]['diff']; put(ws,r,2+j,cpp(q['est'],q['p']))
ws.row_dimensions[r].height=28; r+=1
put(ws,r,1,'  Spell-level slope, conflicted minus other (rank-sum p for medians)',F,LEFT)
sl={'peer_adj_ret100':SL[('peer_adj_ret100','min3')],'b_AB':SL[('b_AB','min3')],'fee':AF['fee_ratio_pct'],'alt':AF['alt_pct']}
for j,(y,_) in enumerate(OUTS): q=sl[y]; dec=3 if y!='fee' else 4; put(ws,r,2+j,f"{q['diff']:+.{dec}f}{st(q['p'])} ({q['p']:.3f}; {q['prs']:.3f})",F,RIGHTW)
ws.row_dimensions[r].height=28; r+=1
ws.cell(row=r,column=1,value='Panel C. Mediation (Mediation sheet)').font=FB; r+=1
for fk,lab in [('fye','  Change in Conflicted when Fee Ratio, Fee Slope and alternatives share are added, FYE FE'),('plan','  Same, plan + FYE FE')]:
    put(ws,r,1,lab,F,LEFT)
    for j,(y,_) in enumerate(YS):
        q=CH[(y,fk,'all')]; b0=MED[(y,fk,'t3')]['coef']['conf_pen_vendor'][0]; b1=MED[(y,fk,'all')]['coef']['conf_pen_vendor'][0]
        put(ws,r,2+j,f"{b0:+.3f} to {b1:+.3f}\nchange {q['diff']:+.3f} (p {q['p']:.2f})",F,RIGHTW)
    ws.row_dimensions[r].height=30; r+=1
rule(ws,r-1,5)
ws.page_setup.orientation='landscape'; ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
out='/home/user/PensionData/referee_response/time_to_departure/mechanism_mediation.xlsx'; wb.save(out); print('saved',out)
