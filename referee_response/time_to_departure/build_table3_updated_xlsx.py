import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
R = pd.read_pickle('table3_newbench_res.pkl'); RES=R['RES']; Y=R['Y']
TNR='Times New Roman'; SZ=10.5
F=Font(name=TNR,size=SZ); FB=Font(name=TNR,size=SZ,bold=True); FI=Font(name=TNR,size=SZ,italic=True); BLUE=Font(name=TNR,size=SZ,color='0000FF')
thin=Side(style='thin'); CENTER=Alignment(horizontal='center',vertical='center',wrap_text=True); RIGHT=Alignment(horizontal='right'); JUST=Alignment(horizontal='justify',wrap_text=True,vertical='top')
LAB={'conf':'Conflicted CIO (0/1)','l_invret':'Investment Return (t - 1)','size':'Log (Plan Size) (t-1)','l_fund':'Funding Ratio (t-1)','l_eq':'Equity Allocation (%)','l_fi':'Fixed Income Allocation (%)','l_re':'Real Estate Allocation (%)','l_pe':'Private Equity Allocation (%)','l_hf':'Hedge Fund Allocation (%)','l_alt':'Misc. Alternatives Allocation (%)','l_comd':'Commodities Allocation (%)','l_cash':'Cash Allocation (%)','fincenter1':'Financial Center (0/1)','l_log_cio_age':'Log (CIO Age) (t-1)','l_log_cio_tenure':'Log (CIO Tenure) (t-1)','prior_private_exp':'CIO Prior Private Experience (0/1)','Elite_inst':'CIO Attended Elite College (0/1)','l_sib1':'Separate Investment Board (0/1)','wstate_political':'% Political Appointees on Board'}
VARS=list(LAB)
wb=Workbook(); E=wb.active; E.title='Estimates'
E['A1']='Estimates behind the tables. Blue = regression output (coefficient, SE, minimum cluster count). t and p are formulas; p uses the t distribution with (min clusters - 1) df. SE double clustered by plan and fiscal year (Cameron-Gelbach-Miller, with the eigenvalue fix).'; E['A1'].font=FI
for j,h in enumerate(['Version','Column','Variable','Coefficient','SE','Clusters (min)','t','p'],1): E.cell(row=3,column=j,value=h).font=FB
E.column_dimensions['A'].width=14; E.column_dimensions['C'].width=36
er=[4]
def reg(ver,col,var,b,se,G):
    r=er[0]; er[0]+=1
    for j,v in enumerate([ver,col,var],1): E.cell(row=r,column=j,value=v).font=F
    E.cell(row=r,column=4,value=b).font=BLUE; E.cell(row=r,column=5,value=se).font=BLUE; E.cell(row=r,column=6,value=G).font=BLUE
    E.cell(row=r,column=7,value=f'=IF(E{r}>0,D{r}/E{r},"")').font=F; E.cell(row=r,column=8,value=f'=IF(G{r}="","",TDIST(ABS(G{r}),F{r}-1,2))').font=F
    for j in [4,5,7,8]: E.cell(row=r,column=j).number_format='0.0000'
    return dict(coef=f'Estimates!D{r}',t=f'Estimates!G{r}',p=f'Estimates!H{r}')
def cf(ref): return f'=IF({ref["p"]}="","n/a",TEXT({ref["coef"]},"0.000")&IF({ref["p"]}<0.01,"***",IF({ref["p"]}<0.05,"**",IF({ref["p"]}<0.1,"*",""))))'
def tf(ref): return f'=IF({ref["t"]}="","","("&TEXT({ref["t"]},"0.00")&")")'
def panel(ws,r,title,deplab,key):
    regs=RES[key]; n=len(regs)
    ws.cell(row=r,column=1,value=title).font=FB; r+=1
    for j in range(n): c=ws.cell(row=r,column=2+j,value=f'({j+1})'); c.font=F; c.alignment=CENTER; c.border=Border(top=thin)
    ws.cell(row=r,column=1).border=Border(top=thin); r+=1
    for j in range(n): c=ws.cell(row=r,column=2+j,value=deplab); c.font=F; c.alignment=CENTER; c.border=Border(bottom=thin)
    ws.cell(row=r,column=1).border=Border(bottom=thin); r+=1
    for v in VARS:
        if not any(v in x['est'] for x in regs): continue
        ws.cell(row=r,column=1,value=LAB[v]).font=F
        for j,x in enumerate(regs):
            if v in x['est']:
                ref=reg(key,f'({j+1})',LAB[v],*x['est'][v],x['G'])
                a=ws.cell(row=r,column=2+j,value=cf(ref)); a.font=F; a.alignment=RIGHT; b=ws.cell(row=r+1,column=2+j,value=tf(ref)); b.font=F; b.alignment=RIGHT
        r+=2
    r+=1
    for lab,vals,fmt in [('Observations',[x['N'] for x in regs],'#,##0'),('Adjusted R-squared',[x['adjr2'] for x in regs],'0.000')]:
        ws.cell(row=r,column=1,value=lab).font=F
        for j,v in enumerate(vals): c=ws.cell(row=r,column=2+j,value=v); c.font=F; c.alignment=RIGHT; c.number_format=fmt
        r+=1
    for lab,k in [('Plan FE','plan'),('Year FE','year'),('Last n years','lastn')]:
        ws.cell(row=r,column=1,value=lab).font=F
        for j,x in enumerate(regs): c=ws.cell(row=r,column=2+j,value=x['col'][k]); c.font=F; c.alignment=CENTER
        r+=1
    for c in range(1,2+n): ws.cell(row=r-1,column=c).border=Border(bottom=thin)
    return r+2
notes=("This table reports results of tests that regress the returns of public pension plans on an indicator variable, Conflicted CIO, which is equal to 1 for all years prior to the plan's investment officer (CIO) departure for a position in a private asset management firm and 0 otherwise. The performance measure in Panel A is Peer-Adjusted Return, which is the difference of a plan's return and the average pension plan return each year. "
 "The performance measure in Panel B is Benchmark-Adjusted Return, which is calculated as the plan's return minus the return of its policy benchmark, when available. Policy benchmark returns are collected from plans' financial reports; where a report does not state the one-year benchmark return, it is inferred from the plan's reported multi-year benchmark returns or from its policy weights and component index returns (estimates graded A or B only). "
 "The other independent variables are as defined in Table I. Columns 5 and 6 redefine Conflicted CIO to equal 1 only in the last 5 or 10 years before departure. The standard errors are double clustered by plan and year. Coefficients marked with ***, **, and * are significant at the 1%, 5%, and 10% level, respectively.")
ws=wb.create_sheet('Table III (updated)'); ws.sheet_view.showGridLines=False; ws.column_dimensions['A'].width=36
for c in range(2,8): ws.column_dimensions[L(c)].width=15
ws.cell(row=1,column=1,value='Table III: Conflicted CIO and Plan Performance Regressions (updated policy benchmark)').font=FB
ws.merge_cells('A2:G2'); ws['A2']=notes; ws['A2'].font=F; ws['A2'].alignment=JUST; ws.row_dimensions[2].height=150
r=4; r=panel(ws,r,'Panel A. Peer-Adjusted Returns','Peer-adj. Ret. (t)','peer'); r=panel(ws,r,'Panel B. Benchmark-Adjusted Returns (updated benchmark)','Bench-adj. Ret. (t)','b_AB')
# comparison sheet
ws=wb.create_sheet('Benchmark versions'); ws.sheet_view.showGridLines=False; ws.column_dimensions['A'].width=70
for c in range(2,8): ws.column_dimensions[L(c)].width=15
ws.cell(row=1,column=1,value='Conflicted CIO coefficient in Table III under alternative benchmark-adjusted returns (same specifications as Table III columns 1-6)').font=FB
n2=("Each row re-estimates the six Table III columns with a different dependent variable; only the Conflicted CIO coefficient is shown, with all controls, fixed effects and clustering as in Table III. 'PPD return' is the plan's investment return in the Public Plans Database (the return used in the paper). 'Document return' is the total-fund return reported on the same page of the plan's financial report as the policy benchmark. "
 "'Reported' benchmarks are stated in the report; 'inferred A/B' and 'all inferred' add benchmark returns inferred from multi-year benchmark returns or from policy weights and component index returns, by estimate grade. The row marked 'agree within 0.5 pp' keeps only plan-years in which the PPD and document returns differ by at most 0.5 percentage points, removing fiscal-year misalignments between the two sources.")
ws.merge_cells('A2:G2'); ws['A2']=n2; ws['A2'].font=F; ws['A2'].alignment=JUST; ws.row_dimensions[2].height=95
r=4
for j in range(6): c=ws.cell(row=r,column=2+j,value=f'({j+1})'); c.font=F; c.alignment=CENTER; c.border=Border(top=thin,bottom=thin)
ws.cell(row=r,column=1,value='Dependent variable').font=F; ws.cell(row=r,column=1).border=Border(top=thin,bottom=thin); r+=1
order=['peer','b_orig','b_rep','b_AB','b_all','b_AB_ok','b_docAB','b_docall']
for k in order:
    ws.cell(row=r,column=1,value=Y[k][1]).font=FB if k in ('b_orig','b_AB') else F
    for j,x in enumerate(RES[k]):
        ref=reg(f'{k} (summary)',f'({j+1})','Conflicted CIO',*x['est']['conf'],x['G'])
        a=ws.cell(row=r,column=2+j,value=cf(ref)); a.font=F; a.alignment=RIGHT; b=ws.cell(row=r+1,column=2+j,value=tf(ref)); b.font=F; b.alignment=RIGHT
        c=ws.cell(row=r+2,column=2+j,value=x['N']); c.font=FI; c.alignment=RIGHT; c.number_format='#,##0'
    ws.cell(row=r+2,column=1,value='   Observations').font=FI
    r+=3
for lab,vals in [('Plan FE',[x['col']['plan'] for x in RES['peer']]),('Year FE',[x['col']['year'] for x in RES['peer']]),('Last n years',[x['col']['lastn'] for x in RES['peer']])]:
    ws.cell(row=r,column=1,value=lab).font=F
    for j,v in enumerate(vals): c=ws.cell(row=r,column=2+j,value=v); c.font=F; c.alignment=CENTER
    r+=1
for c in range(1,8): ws.cell(row=r-1,column=c).border=Border(bottom=thin)
r+=1
ws.cell(row=r,column=1,value='Note on t-statistics: coefficients and observations match reghdfe exactly. In the year-FE-only columns (1), (2), (5) and (6), these two-way clustered t-statistics run about 4 to 17 percent larger than reghdfe reports for the published Table III (for example, original Panel B column 1: -2.65 here, -2.27 published); plan-FE columns (3) and (4) match to within 0.06. Use the attached Stata code for the exact figures.').font=FI
ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=7); ws.cell(row=r,column=1).alignment=JUST; ws.row_dimensions[r].height=45
wb.move_sheet('Estimates', offset=2)
out='/home/user/PensionData/referee_response/time_to_departure/table3_updated_benchmark.xlsx'; wb.save(out); print('saved')
