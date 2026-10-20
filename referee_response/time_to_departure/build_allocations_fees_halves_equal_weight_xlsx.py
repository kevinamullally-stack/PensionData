import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
D=pd.read_pickle('alloc_fee_halves_ew.pkl'); R=D['R']; SL=D['SL']; YS=D['YS']; RATE=D['RATE']; REG=D['REG']; C3=D['coef3']
TNR='Times New Roman'; SZ=10.5
F=Font(name=TNR,size=SZ); FB=Font(name=TNR,size=SZ,bold=True); FI=Font(name=TNR,size=SZ,italic=True)
thin=Side(style='thin'); CENTER=Alignment(horizontal='center',vertical='center',wrap_text=True); RIGHT=Alignment(horizontal='right',vertical='center'); LEFT=Alignment(horizontal='left',vertical='center',wrap_text=True); JUST=Alignment(horizontal='justify',wrap_text=True,vertical='top')
def st(p): return '' if p is None or np.isnan(p) else '***' if p<0.01 else '**' if p<0.05 else '*' if p<0.1 else ''
def cp(v,p,dec=3): return f"{v:+.{dec}f}{st(p)} ({p:.3f})"
ALLOC=[y for y,_ in YS if y.endswith('_pct') and not y.startswith('fee')]; FEES=[y for y,_ in YS if y.startswith('fee')]
LAB=dict(YS); DEC={y:(2 if y in ALLOC else 3) for y in LAB}
FEENOTE=(f"Allocation-adjusted Fee Ratios are residuals from pooled OLS regressions of the Fee Ratio on the asset-class shares over the Table III sample: on the equity, fixed income and alternatives shares "
         f"(coefficients {C3['eq_pct']:+.4f}, {C3['fi_pct']:+.4f} and {C3['alt_pct']:+.4f} per percentage point), on the same three shares and fiscal-year-end date effects, and on all seven asset-class shares "
         "(equity, fixed income, real estate, private equity, hedge funds, other alternatives, commodities; cash omitted) and fiscal-year-end date effects. ")
BASENOTE=("Sample: CIOs who have left (curr = 0) in the Table III sample with defined halves of tenure. For each CIO-plan spell (or each CIO, averaging across plans within a CIO-year), the change is the mean of the outcome in the second half of tenure minus its mean in the first half; every spell (CIO) counts once. "
          "The second half is the years in which tenure exceeds half of tenure at departure. Allocations are percentages of plan assets (Alternatives = real estate + private equity + hedge funds + other alternatives + commodities); Fee Ratio = 100 x total_fee_pct. ")
STATNOTE=("Columns report the mean change for conflicted CIOs with the p-value of a test that it is zero, their median change with the signed-rank p-value, the mean change for other departing CIOs, the conflicted-minus-other difference with p-values from a t-test, a t-test allowing unequal variances (Welch), and a rank-sum test, and the difference when spells are weighted by their plan-years (standard errors clustered by plan). "
          "'Net of fiscal-year-end mean' subtracts the mean of all Table III plan-years with the same fiscal-year-end date (for residuals that already net out fiscal-year-end effects this row equals the first row). 'Drop' rows remove the conflicted spells with the most negative or most positive changes. "
          "The last row of each panel repeats the comparison with each spell's OLS slope of the outcome on fiscal year (at least three years, winsorized at 1% and 99%), parallel to the Fee Slope. Entries are estimates with p-values in parentheses; ***, **, and * denote significance at the 1%, 5%, and 10% level.")
HEAD=['','Conflicted: mean change (p vs. 0)','Conflicted: median change (signed-rank p)','Other departing CIOs: mean change (p vs. 0)','Conflicted minus other (t-test p)','Welch p','Rank-sum p','Weighted by plan-years (p)','Spells (conflicted / other)']
ROWS=[('spell','base','Each CIO-plan spell'),('spell','ten4','  Tenure at departure of at least 4 years'),('spell','dm','  Net of fiscal-year-end mean'),
      ('spell','dropneg3','  Drop the 3 most negative conflicted changes'),('spell','droppos3','  Drop the 3 most positive conflicted changes'),
      ('spell','dropneg5','  Drop the 5 most negative conflicted changes'),('spell','droppos5','  Drop the 5 most positive conflicted changes'),
      ('cio','base','One observation per CIO'),('cio','ten4','  Tenure at departure of at least 4 years'),('cio','dm','  Net of fiscal-year-end mean')]
def setup(ws,title,note,W):
    ws.sheet_view.showGridLines=False
    for j,w in enumerate(W): ws.column_dimensions[L(1+j)].width=w
    ws['A1']=title; ws['A1'].font=FB
    ws.merge_cells(f'A2:{L(len(W))}2'); ws['A2']=note; ws['A2'].font=F; ws['A2'].alignment=JUST; ws.row_dimensions[2].height=14*(len(note)//(sum(W)*1.3)+1)
    ws.page_setup.orientation='landscape'; ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
def header(ws,r,H,h=58):
    for j,x_ in enumerate(H): x=ws.cell(row=r,column=1+j,value=x_); x.font=F; x.alignment=CENTER; x.border=Border(top=thin,bottom=thin)
    ws.row_dimensions[r].height=h
def rowvals(ws,r,lab,vals,font=None):
    ws.cell(row=r,column=1,value=lab).font=font or (F if lab.startswith('  ') else FB); ws.cell(row=r,column=1).alignment=LEFT
    for j,v in enumerate(vals): x=ws.cell(row=r,column=2+j,value=v); x.font=font or F; x.alignment=RIGHT
def detail(ws,r,ys,first_panel='A'):
    for i,y in enumerate(ys):
        ws.cell(row=r,column=1,value=f"Panel {chr(ord(first_panel)+i)}. {LAB[y]}").font=FB; r+=1
        header(ws,r,HEAD); r+=1; dec=DEC[y]
        for unit,k,lab in ROWS:
            q=R[(y,unit,k)]
            rowvals(ws,r,lab,[cp(q['mc'],q['p1c'],dec),cp(q['medc'],q['psr'],dec),cp(q['mo'],q['p1o'],dec),cp(q['diff'],q['p'],dec),f"{q['pw']:.3f}",f"{q['prs']:.3f}",cp(q['wdiff'],q['wp'],dec),f"{q['nc']} / {q['no']}"]); r+=1
        q=SL[y]; sd=dec+1
        rowvals(ws,r,'Spell-level slope (pp per year)',[cp(q['mc'],q['p1c'],sd),f"{q['medc']:+.{sd}f}",cp(q['mo'],q['p1o'],sd),cp(q['diff'],q['p'],sd),f"{q['pw']:.3f}",f"{q['prs']:.3f}",cp(q['wdiff'],q['wp'],sd),f"{q['nc']} / {q['no']}"],FI); r+=1
        for c in range(1,10): ws.cell(row=r-1,column=c).border=Border(bottom=thin)
        r+=1
    return r
wb=Workbook()
# ---------------- Summary ----------------
ws=wb.active; ws.title='Summary'
setup(ws,'Change in Allocations and Allocation-Adjusted Fees From the First to the Second Half of Tenure, Each CIO Weighted Equally: Summary',
      BASENOTE+FEENOTE+"Each row reports, for CIO-plan spells, the mean change for conflicted and other departing CIOs and the conflicted-minus-other difference (t-test p-value; Welch and rank-sum p-values), the difference net of the fiscal-year-end mean, the difference with one observation per CIO, and the difference in spell-level slopes (pp per year). Full detail is on the next two sheets. ***, **, and * denote significance at the 1%, 5%, and 10% level.",
      [52,14,14,18,10,10,18,18,18])
r=4
H=['','Conflicted: mean change','Other: mean change','Conflicted minus other (p)','Welch p','Rank-sum p','Net of fiscal-year-end mean: difference (p)','One observation per CIO: difference (p)','Spell-level slope: difference (p)']
for title,ys in [('Panel A. Allocations (% of plan assets)',ALLOC),('Panel B. Fee Ratio and allocation-adjusted Fee Ratio (%)',FEES)]:
    ws.cell(row=r,column=1,value=title).font=FB; r+=1; header(ws,r,H,44); r+=1
    for y in ys:
        q=R[(y,'spell','base')]; d=DEC[y]; qd=R[(y,'spell','dm')]; qc=R[(y,'cio','base')]; qs=SL[y]
        dmtxt='same as base' if y in ('fee_r3fe','fee_r7fe') else cp(qd['diff'],qd['p'],d)
        rowvals(ws,r,'  '+LAB[y],[f"{q['mc']:+.{d}f}",f"{q['mo']:+.{d}f}",cp(q['diff'],q['p'],d),f"{q['pw']:.3f}",f"{q['prs']:.3f}",dmtxt,cp(qc['diff'],qc['p'],d),cp(qs['diff'],qs['p'],d+1)]); r+=1
    for c in range(1,10): ws.cell(row=r-1,column=c).border=Border(bottom=thin)
    r+=1
# ---------------- Allocations ----------------
ws=wb.create_sheet('Allocations')
setup(ws,'Change in Asset Allocations From the First to the Second Half of Tenure, Each CIO Weighted Equally',BASENOTE+STATNOTE,[44,18,20,18,18,10,10,18,12])
detail(ws,4,ALLOC)
# ---------------- Fees ----------------
ws=wb.create_sheet('Fee Ratio net of allocations')
setup(ws,'Change in the Fee Ratio, Before and After Controlling for Asset Allocations, From the First to the Second Half of Tenure, Each CIO Weighted Equally',
      BASENOTE+FEENOTE+STATNOTE+" Panel E divides each spell's change by the number of fiscal years between the midpoints of its two halves, and regresses each spell's change (or slope) on Conflicted and the spell's own change (or slope) in the equity, fixed income and alternatives shares, equal weights, standard errors clustered by plan.",
      [44,18,20,18,18,10,10,18,12])
r=detail(ws,4,FEES)
ws.cell(row=r,column=1,value='Panel E. Per-year change and controls for each spell\'s own allocation changes').font=FB; r+=1
header(ws,r,['','Conflicted: mean (p vs. 0)','Conflicted: median (signed-rank p)','Other departing CIOs: mean (p vs. 0)','Conflicted minus other (t-test p)','Welch p','Rank-sum p','Weighted by plan-years (p)','Spells (conflicted / other)']); r+=1
for y in FEES:
    ws.cell(row=r,column=1,value=LAB[y]).font=FB; r+=1
    q=RATE[y]
    rowvals(ws,r,f"  Change per year between half midpoints (years apart: {q['gapc']:.2f} / {q['gapo']:.2f})",[cp(q['mc'],q['p1c'],4),cp(q['medc'],q['psr'],4),cp(q['mo'],q['p1o'],4),cp(q['diff'],q['p'],4),f"{q['pw']:.3f}",f"{q['prs']:.3f}",cp(q['wdiff'],q['wp'],4),f"{q['nc']} / {q['no']}"]); r+=1
    for k,lab in [('halves','  Halves change on Conflicted + own changes in the three shares'),('slope','  Spell slope on Conflicted + own slopes of the three shares')]:
        q=REG[(k and (y,k))]
        rowvals(ws,r,lab,['','','',cp(q['b'],q['p'],4),'','','',f"N = {q['n']}"]); r+=1
for c in range(1,10): ws.cell(row=r-1,column=c).border=Border(bottom=thin)
out='/home/user/PensionData/referee_response/time_to_departure/allocations_fees_halves_equal_weight.xlsx'; wb.save(out); print('saved',out)
