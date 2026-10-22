import pandas as pd, numpy as np
from scipy import stats
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
D=pd.read_pickle('table3_conflicted_halves_preshock.pkl'); RES0=D['RES']; YS=D['YS']; CNT=D['CNT']; PERL=D['PER']
TNR='Times New Roman'; SZ=10.5
F=Font(name=TNR,size=SZ); FB=Font(name=TNR,size=SZ,bold=True); FI=Font(name=TNR,size=SZ,italic=True)
thin=Side(style='thin'); CENTER=Alignment(horizontal='center',vertical='center',wrap_text=True); RIGHT=Alignment(horizontal='right'); LEFT=Alignment(horizontal='left',vertical='center',wrap_text=True); JUST=Alignment(horizontal='justify',wrap_text=True,vertical='top')
def st(p): return '***' if p<0.01 else '**' if p<0.05 else '*' if p<0.1 else ''
LAB={'h1':'Conflicted CIO x First Half of Tenure (0/1)','h2':'Conflicted CIO x Second Half of Tenure (0/1)',
 'l_invret':'Investment Return (t-1)','size':'Log (Plan Size) (t-1)','l_fund':'Funding Ratio (t-1)','l_eq':'Equity Allocation (t-1)','l_fi':'Fixed Income Allocation (t-1)','l_re':'Real Estate Allocation (t-1)',
 'l_pe':'Private Equity Allocation (t-1)','l_hf':'Hedge Fund Allocation (t-1)','l_alt':'Misc. Alternatives Allocation (t-1)','l_comd':'Commodities Allocation (t-1)','l_cash':'Cash Allocation (t-1)','fincenter1':'Financial Center (0/1)',
 'l_log_cio_age':'Log (CIO Age) (t-1)','l_log_cio_tenure':'Log (CIO Tenure) (t-1)','prior_private_exp':'CIO Prior Private Experience (0/1)','Elite_inst':'CIO Attended Elite College (0/1)','l_sib1':'Separate Investment Board (0/1)','wstate_political':'% Political Appointees on Board'}
def base(per):
    c=CNT[per]
    span=("fiscal years 2002-2008 (the Table III sample through FY2008, before the 2009-2014 shock period)" if per=='pre2009' else "the Table III sample excluding fiscal years 2009-2014 (the shock period)")
    return ("This table repeats Table III with fiscal-year-end date fixed effects, with Conflicted CIO split into two indicators, estimated on "+span+". "
      "Conflicted CIO x First Half of Tenure equals 1 in a conflicted CIO's plan-years in which tenure at the plan is at most half of tenure at departure, and Conflicted CIO x Second Half of Tenure equals 1 in the remaining conflicted plan-years; halves are defined from each CIO's full tenure, so only the estimation years change "
      f"({c['h1']} first-half and {c['h2']} second-half conflicted plan-years in this period). In columns 1-3 the reference group is the non-conflicted plan-years. Column 4 uses only conflicted CIOs' plan-years and adds CIO fixed effects, so First Half is the reference and the Second Half coefficient is the within-CIO change from the first to the second half. "
      "The performance measure in Panel A is Peer-Adjusted Return; in Panel B it is Benchmark-Adjusted Return (plan return minus its policy benchmark: reported, else inferred grade A or B). CIO fixed effects are for each CIO's tenure at a plan; Financial Center, CIO Prior Private Experience and CIO Attended Elite College do not vary within a CIO's tenure and are absorbed. "
      "'Second minus first half' is the difference between the two half coefficients. t-statistics in parentheses; standard errors double clustered by plan and year"+(" (only seven year clusters, so p-values with standard errors clustered by plan alone are also reported)" if per=='pre2009' else " (p-values with standard errors clustered by plan alone also reported)")+". ***, **, and * denote significance at the 1%, 5%, and 10% level.")
SHEETS=[(per,v,nm,PERL[per]+(': Without Log CIO Tenure and Log CIO Age' if v=='paper_noten' else ''),(' Log (CIO Tenure) and Log (CIO Age) are dropped from the controls because tenure rises mechanically from the first to the second half.' if v=='paper_noten' else ''))
        for per,v,nm in [('pre2009','paper','Through FY2008'),('pre2009','paper_noten','Through FY2008, no tenure'),('ex0914','paper','Excl. FY2009-2014'),('ex0914','paper_noten','Excl. FY2009-2014, no tenure')]]
NCOL=4
wb=Workbook(); first=True
for per,ver,sname,title,extra in SHEETS:
    RES={(ver,y):RES0[(per,ver,y)] for y,_ in YS}
    ws=wb.active if first else wb.create_sheet(); first=False; ws.title=sname; ws.sheet_view.showGridLines=False
    ws.column_dimensions['A'].width=46
    for j in range(2,2+NCOL): ws.column_dimensions[L(j)].width=17
    ws['A1']='Table III with Conflicted CIO Split Into the First and Second Half of Tenure: '+title; ws['A1'].font=FB
    note=base(per)+extra; ws.merge_cells(f'A2:{L(1+NCOL)}2'); ws['A2']=note; ws['A2'].font=F; ws['A2'].alignment=JUST; ws.row_dimensions[2].height=14*(len(note)//118+1)
    r=4
    for pi,(y,yl) in enumerate(YS):
        regs=RES[(ver,y)]
        ws.cell(row=r,column=1,value=f"Panel {'AB'[pi]}. {yl}").font=FB; r+=1
        for j in range(NCOL): x=ws.cell(row=r,column=2+j,value=f'({j+1})'); x.font=F; x.alignment=CENTER; x.border=Border(top=thin)
        ws.cell(row=r,column=1).border=Border(top=thin); r+=1
        for j in range(NCOL): x=ws.cell(row=r,column=2+j,value=('Peer-adj. Ret. (t)' if pi==0 else 'Bench-adj. Ret. (t)')); x.font=F; x.alignment=CENTER; x.border=Border(bottom=thin)
        ws.cell(row=r,column=1).border=Border(bottom=thin); r+=1
        for v,l in LAB.items():
            if not any(v in x['est'] for x in regs): continue
            bold=v in ('h1','h2')
            ws.cell(row=r,column=1,value=l).font=FB if bold else F; ws.cell(row=r,column=1).alignment=LEFT
            for j,x in enumerate(regs):
                if v=='h1' and x['within']:
                    a=ws.cell(row=r,column=2+j,value='Reference'); a.font=F; a.alignment=RIGHT; continue
                if v not in x['est']: continue
                b,se=x['est'][v]; t=b/se; p=2*stats.t.sf(abs(t),x['G']-1)
                a=ws.cell(row=r,column=2+j,value=f"{b:.3f}{st(p)}"); a.font=FB if bold else F; a.alignment=RIGHT
                a=ws.cell(row=r+1,column=2+j,value=f"({t:.2f})"); a.font=F; a.alignment=RIGHT
            r+=2
        ws.cell(row=r,column=1,value='Second minus first half').font=FB
        ws.cell(row=r+1,column=1,value='  t-statistic').font=F; ws.cell(row=r+2,column=1,value='  p-value').font=F; ws.cell(row=r+3,column=1,value='  p-value, clustered by plan').font=F
        for j,x in enumerate(regs):
            b,se=x['diff']; t=b/se; p=2*stats.t.sf(abs(t),x['G']-1)
            a=ws.cell(row=r,column=2+j,value=f"{b:.3f}{st(p)}"); a.font=FB; a.alignment=RIGHT
            a=ws.cell(row=r+1,column=2+j,value=f"({t:.2f})"); a.font=F; a.alignment=RIGHT
            a=ws.cell(row=r+2,column=2+j,value=f"{p:.3f}"); a.font=F; a.alignment=RIGHT
            a=ws.cell(row=r+3,column=2+j,value=f"{x['pp']['diff']:.3f}"); a.font=F; a.alignment=RIGHT
        r+=5
        for lab,vals in [('Sample',['Conflicted CIOs' if x['within'] else 'All' for x in regs]),('Observations',[f"{x['N']:,}" for x in regs]),('Year clusters',[f"{x['Gy']}" for x in regs]),('Adjusted R-squared',[f"{x['adjr2']:.3f}" for x in regs]),
                         ('Plan FE',[x['col']['plan'] for x in regs]),('CIO FE',[x['col']['cio'] for x in regs]),('Fiscal-year-end FE',[x['col']['year'] for x in regs])]:
            ws.cell(row=r,column=1,value=lab).font=F
            for j,v in enumerate(vals): a=ws.cell(row=r,column=2+j,value=v); a.font=F; a.alignment=RIGHT if lab in ('Observations','Adjusted R-squared','Year clusters') else CENTER
            r+=1
        for cc in range(1,2+NCOL): ws.cell(row=r-1,column=cc).border=Border(bottom=thin)
        r+=2
    ws.page_setup.orientation='portrait'; ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
out='/home/user/PensionData/referee_response/time_to_departure/table3_conflicted_halves_preshock.xlsx'; wb.save(out); print('saved',out)
