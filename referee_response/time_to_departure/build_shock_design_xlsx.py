import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter as L
D=pd.read_pickle('shock_tests.pkl'); R=D['R']; YS=D['YS']; FES=D['FES']; EVB=D['EVB']; BINS=D['BINS']; EXPO=D['EXPO']; STARTS=D['STARTS']; GRP=D['GRP']
TNR='Times New Roman'; SZ=10.5
F=Font(name=TNR,size=SZ); FB=Font(name=TNR,size=SZ,bold=True); FI=Font(name=TNR,size=SZ,italic=True); BLUE=Font(name=TNR,size=SZ,color='0000FF')
thin=Side(style='thin'); CENTER=Alignment(horizontal='center',vertical='center',wrap_text=True); RIGHT=Alignment(horizontal='right'); LEFT=Alignment(horizontal='left',vertical='center',wrap_text=True); JUST=Alignment(horizontal='justify',wrap_text=True,vertical='top')
SHADE=PatternFill('solid',fgColor='F1F0EC')
YL={'peer_adj_ret100':'Peer-Adjusted Return','b_AB':'Benchmark-Adjusted Return'}; FL={'fy':'Year FE','ppd_id':'Plan FE','ppd_id + fy':'Plan and Year FE'}
wb=Workbook(); E=wb.active; E.title='Estimates'
E['A1']='Estimates behind the tables. Blue = regression output (coefficient, SE, minimum cluster count). t and p are formulas; p uses the t distribution with (min clusters - 1) df. SE double clustered by plan and year, with the eigenvalue fix.'; E['A1'].font=FI
for j,h in enumerate(['Table','Dependent variable','Specification','Term','Coefficient','SE','Clusters (min)','t','p'],1): E.cell(row=3,column=j,value=h).font=FB
for c,w in zip('ABCD',[22,28,34,34]): E.column_dimensions[c].width=w
er=[4]
def reg(tab,y,spec,term,r,G):
    i=er[0]; er[0]+=1
    for j,v in enumerate([tab,YL[y],spec,term],1): E.cell(row=i,column=j,value=v).font=F
    E.cell(row=i,column=5,value=None if np.isnan(r['est']) else r['est']).font=BLUE; E.cell(row=i,column=6,value=None if np.isnan(r['se']) else r['se']).font=BLUE; E.cell(row=i,column=7,value=G).font=BLUE
    E.cell(row=i,column=8,value=f'=IF(AND(ISNUMBER(E{i}),F{i}>0),E{i}/F{i},"")').font=F; E.cell(row=i,column=9,value=f'=IF(H{i}="","",TDIST(ABS(H{i}),G{i}-1,2))').font=F
    for j in [5,6,8,9]: E.cell(row=i,column=j).number_format='0.0000'
    return dict(coef=f'Estimates!E{i}',t=f'Estimates!H{i}',p=f'Estimates!I{i}')
def cf(ref): return f'=IF({ref["p"]}="","",TEXT({ref["coef"]},"0.000")&IF({ref["p"]}<0.01,"***",IF({ref["p"]}<0.05,"**",IF({ref["p"]}<0.1,"*",""))))'
def tf(ref): return f'=IF({ref["t"]}="","","("&TEXT({ref["t"]},"0.00")&")")'
def put(ws,r,c,ref,fnt=F):
    a=ws.cell(row=r,column=c,value=cf(ref)); a.font=fnt; a.alignment=RIGHT; b=ws.cell(row=r+1,column=c,value=tf(ref)); b.font=fnt; b.alignment=RIGHT
def sheet(name,title,note,ncol,widths,per=110):
    ws=wb.create_sheet(name); ws.sheet_view.showGridLines=False
    for j,w in enumerate(widths): ws.column_dimensions[L(1+j)].width=w
    ws['A1']=title; ws['A1'].font=FB
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=ncol); ws['A2']=note; ws['A2'].font=F; ws['A2'].alignment=JUST; ws.row_dimensions[2].height=14*(len(note)//per+2)
    ws.page_setup.orientation='landscape'; ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    return ws
def header(ws,r,groups,sub):     # groups: list of (label, ncols); sub: list of subheaders
    c=2
    for lab,n in groups:
        if n>1: ws.merge_cells(start_row=r,start_column=c,end_row=r,end_column=c+n-1)
        x=ws.cell(row=r,column=c,value=lab); x.font=F; x.alignment=CENTER
        for k in range(n): ws.cell(row=r,column=c+k).border=Border(top=thin)
        c+=n
    ws.cell(row=r,column=1).border=Border(top=thin); ws.row_dimensions[r].height=30
    for j,h in enumerate(sub): x=ws.cell(row=r+1,column=2+j,value=h); x.font=F; x.alignment=CENTER; x.border=Border(bottom=thin)
    ws.cell(row=r+1,column=1).border=Border(bottom=thin); ws.row_dimensions[r+1].height=30
    return r+2
def stats_row(ws,r,lab,vals,fmt,fnt=F):
    ws.cell(row=r,column=1,value=lab).font=F
    for j,v in enumerate(vals):
        if v is None: continue
        x=ws.cell(row=r,column=2+j,value=v); x.font=fnt; x.alignment=RIGHT; x.number_format=fmt
    return r+1
def bottom(ws,r,ncol):
    for c in range(1,ncol+1): ws.cell(row=r-1,column=c).border=Border(bottom=thin)
CTXT="The control variables are those of Table VIII (the lagged return, Size, the lagged funding ratio, lagged allocations, Financial Center, lagged log CIO age and tenure, lagged sibling plans and the political-state indicator). "
BTXT="Benchmark-Adjusted Return is the plan's return minus the updated policy benchmark (the benchmark return reported in the plan's financial report or, where none is stated, an inferred value graded A or B). "
SETXT="The standard errors are double clustered by plan and year. Coefficients marked with ***, **, and * are significant at the 1%, 5%, and 10% level, respectively."
COLS=[(y,fe) for y,_ in YS for fe,_ in FES]
# ---- Table VIII ----
ws=sheet('Table VIII','Table VIII: Conflicted CIOs and the Loglisci Shock, Peer- and Benchmark-Adjusted Returns',
 "Shock Period equals 1 for plan fiscal years ending between March 2009 and March 2014. The first row reports Conflicted CIO from the same regression without the Shock Period terms (the average effect of conflict over the whole sample). The next rows add Shock Period and Conflicted CIO x Shock Period: Conflicted CIO is then the effect of conflict outside the Shock Period, and the italic row is the effect within it (Conflicted CIO + Conflicted CIO x Shock Period). "+CTXT+BTXT+SETXT,7,[44]+[13]*6)
r=header(ws,4,[(YL[y],3) for y,_ in YS],[FL[fe] for y,fe in COLS])
for lab,key,term,fnt in [('Conflicted CIO, no Shock Period terms','t3','conf',F),('Conflicted CIO (outside Shock Period)','t8','conf',F),('Conflicted CIO x Shock Period','t8','cxs',F),('Effect of conflict within Shock Period','t8','in',FI)]:
    ws.cell(row=r,column=1,value=lab).font=fnt
    for j,(y,fe) in enumerate(COLS): x=R[(key,y,fe)]; put(ws,r,2+j,reg('Table VIII',y,FL[fe],lab,x[term],x['G']),fnt)
    r+=2 if term!='conf' or key!='t3' else 3
r+=1
r=stats_row(ws,r,'Observations',[R[('t8',y,fe)]['N'] for y,fe in COLS],'#,##0'); r=stats_row(ws,r,'Adjusted R-squared',[R[('t8',y,fe)]['adjr2'] for y,fe in COLS],'0.000'); bottom(ws,r,7)
# ---- Event study FY2008 ----
LAB={'es_pre':'Conflicted CIO x FY<2005','es_post':'Conflicted CIO x FY>2017'}
ws=sheet('Event study FY2008','Event Study Around the Loglisci Shock: FY2008 Base Year',
 "Each regression includes Conflicted CIO and Conflicted CIO interacted with indicators for every fiscal year except FY2008 (the last fiscal year before the Shock Period), with endpoint bins for years before 2005 and after 2017. Conflicted CIO is therefore the effect of conflict in FY2008 and each interaction is the change in that effect relative to FY2008. Shaded rows are the Shock Period. "
 "Tests: pre-shock leads (the pre-2005 bin and FY2005-2007) jointly zero; shock-period coefficients (FY2009-2013) jointly zero; post-period coefficients (FY2014-2017 and the post-2017 bin) jointly zero. The italic rows are the mean FY2009-2013 coefficient, the mean FY2014-2017 coefficient, and their difference (a test that the effect of conflict returns after the case is resolved). "+CTXT+BTXT+SETXT,7,[40]+[13]*6)
r=header(ws,4,[(YL[y],3) for y,_ in YS],[FL[fe] for y,fe in COLS])
ws.cell(row=r,column=1,value='Conflicted CIO (effect in FY2008)').font=F
for j,(y,fe) in enumerate(COLS): x=R[('es',y,fe)]; put(ws,r,2+j,reg('Event FY2008',y,FL[fe],'conf',x['conf'],x['G']))
r+=3
for v in EVB:
    lab=LAB.get(v,'Conflicted CIO x FY'+v[3:]); t=int(v[3:]) if v[3:].isdigit() else None
    ws.cell(row=r,column=1,value=lab).font=F
    for j,(y,fe) in enumerate(COLS): x=R[('es',y,fe)]; put(ws,r,2+j,reg('Event FY2008',y,FL[fe],v,x[v],x['G']))
    if t and 2009<=t<=2013:
        for rr in (r,r+1):
            for c in range(1,8): ws.cell(row=rr,column=c).fill=SHADE
    r+=2
r+=1
for lab,key in [('Mean FY2009-2013 coefficient','win_mean'),('Mean FY2014-2017 coefficient','post_mean'),('Post-period mean minus shock-period mean','return')]:
    ws.cell(row=r,column=1,value=lab).font=FI
    for j,(y,fe) in enumerate(COLS): x=R[('es',y,fe)]; put(ws,r,2+j,reg('Event FY2008',y,FL[fe],lab,x[key],x['G']),FI)
    r+=2
r+=1
for lab,key in [('p-value: pre-shock leads jointly zero','pre'),('p-value: shock period jointly zero','win'),('p-value: post period jointly zero','post')]:
    r=stats_row(ws,r,lab,[R[('es',y,fe)][key]['p'] for y,fe in COLS],'0.000',BLUE)
r=stats_row(ws,r,'Observations',[R[('es',y,fe)]['N'] for y,fe in COLS],'#,##0'); bottom(ws,r,7)
# ---- Event study bins ----
BL={'b0102':'2001-02','b0304':'2003-04','b0506':'2005-06','b0910':'2009-10','b1113':'2011-13','b1415':'2014-15','b1617':'2016-17','b1820':'2018-20'}
ws=sheet('Event study bins','Event Study Around the Loglisci Shock: Two-Year Bins, FY2007-08 Base',
 "As in the FY2008 event study, but with Conflicted CIO interacted with two-year periods (the Shock Period split into FY2009-10 and FY2011-13, and FY2018-20 pooled); FY2007-08 is the omitted base period, so Conflicted CIO is the effect of conflict in FY2007-08. Shaded rows are the Shock Period. "+CTXT+BTXT+SETXT,7,[40]+[13]*6)
r=header(ws,4,[(YL[y],3) for y,_ in YS],[FL[fe] for y,fe in COLS])
ws.cell(row=r,column=1,value='Conflicted CIO (effect in FY2007-08)').font=F
for j,(y,fe) in enumerate(COLS): x=R[('bins',y,fe)]; put(ws,r,2+j,reg('Event bins',y,FL[fe],'conf',x['conf'],x['G']))
r+=3
for b,_,_ in BINS:
    ws.cell(row=r,column=1,value='Conflicted CIO x FY'+BL[b]).font=F
    for j,(y,fe) in enumerate(COLS): x=R[('bins',y,fe)]; put(ws,r,2+j,reg('Event bins',y,FL[fe],b,x[b],x['G']))
    if b in ('b0910','b1113'):
        for rr in (r,r+1):
            for c in range(1,8): ws.cell(row=rr,column=c).fill=SHADE
    r+=2
r+=1
for lab,key in [('p-value: pre-shock bins jointly zero','pre'),('p-value: shock-period bins jointly zero','win'),('p-value: post-period bins jointly zero','post')]:
    r=stats_row(ws,r,lab,[R[('bins',y,fe)][key]['p'] for y,fe in COLS],'0.000',BLUE)
r=stats_row(ws,r,'Observations',[R[('bins',y,fe)]['N'] for y,fe in COLS],'#,##0'); bottom(ws,r,7)
# ---- Exposure ----
C2=[(y,fe) for y,_ in YS for fe in ['fy','ppd_id + fy']]
PEM=D['PEMEAN']; PEW=D['PEMEAN_W']
ws=sheet('Exposure','Differential Exposure: Conflicted CIO x Shock Period With Allocation Controls',
 "Each row re-estimates Table VIII adding interactions that let the effect of asset allocation on returns differ in the Shock Period or in every fiscal year. Rows (2)-(3) interact the lagged private equity share, or all eight lagged allocation shares, with Shock Period; rows (4)-(5) interact them with fiscal-year indicators. The table reports Conflicted CIO x Shock Period. "
 f"The lagged private equity share averages {PEM[1]:.2f}% for conflicted and {PEM[0]:.2f}% for other plan-years ({PEW[1]:.2f}% and {PEW[0]:.2f}% in the Shock Period). "+CTXT+BTXT+SETXT,5,[52]+[15]*4)
r=header(ws,4,[(YL[y],2) for y,_ in YS],[FL[fe] for y,fe in C2])
for i,(k,lab) in enumerate(EXPO):
    ws.cell(row=r,column=1,value=f'({i+1}) {lab}').font=F; ws.cell(row=r,column=1).alignment=LEFT
    for j,(y,fe) in enumerate(C2): x=R[('expo',k,y,fe)]; put(ws,r,2+j,reg('Exposure',y,FL[fe],lab,x['cxs'],x['G']))
    r+=2
r=stats_row(ws,r,'Observations',[R[('expo','base',y,fe)]['N'] for y,fe in C2],'#,##0'); bottom(ws,r,5)
# ---- Placebo windows ----
ws=sheet('Placebo windows','Placebo Windows: Conflicted CIO x Five-Year Window',
 "Panel A re-estimates Table VIII with Shock Period replaced by an indicator for a five-year window of fiscal years (the window's main effect is absorbed by the year fixed effects where they are included); the shaded row is the actual window, FY2009-2013 (which differs from Shock Period only for plans with fiscal years ending in January-March 2014). Windows that overlap FY2009-2013 partly contain the actual shock. "
 "Panel B restricts the sample to FY2008 and earlier and compares FY2005-2008 with FY2001-2004, a window that lies entirely before the shock. "+CTXT+BTXT+SETXT,5,[34]+[15]*4)
r=header(ws,4,[(YL[y],2) for y,_ in YS],[FL[fe] for y,fe in C2])
ws.cell(row=r,column=1,value='Panel A. Conflicted CIO x Window').font=FB; r+=1
for s in STARTS:
    lab=f'FY{s}-{s+4}'+(' (actual)' if s==2009 else '')
    ws.cell(row=r,column=1,value=lab).font=FB if s==2009 else F
    for j,(y,fe) in enumerate(C2): x=R[('roll',s,y,fe)]; put(ws,r,2+j,reg('Placebo windows',y,FL[fe],lab,x['cxw'],x['G']))
    if s==2009:
        for rr in (r,r+1):
            for c in range(1,6): ws.cell(row=rr,column=c).fill=SHADE
    r+=2
r+=1; ws.cell(row=r,column=1,value='Panel B. Pre-shock sample (FY<=2008)').font=FB; r+=1
ws.cell(row=r,column=1,value='Conflicted CIO x FY2005-2008').font=F
for j,(y,fe) in enumerate(C2): x=R[('prefake',y,fe)]; put(ws,r,2+j,reg('Placebo windows',y,FL[fe],'pre-period fake window',x['cxf'],x['G']))
r+=2
r=stats_row(ws,r,'Observations (Panel B)',[R[('prefake',y,fe)]['N'] for y,fe in C2],'#,##0'); bottom(ws,r,5)
# ---- Placebo groups ----
C6=[(y,fe,t) for y,_ in YS for fe in ['fy','ppd_id + fy'] for t in ['cxs','oxs','diff']]
ws=sheet('Placebo groups','Placebo Groups: Other Departures x Shock Period',
 "Each regression adds to Table VIII an indicator for CIOs who later leave for a reason other than a vendor job (and are not conflicted) and its interaction with Shock Period. The omitted group is CIOs who do not leave during the sample. If the shock worked through vendor-employment conflicts, Conflicted CIO x Shock Period should be positive and Other x Shock Period should not; the third column of each block tests their difference. "
 "Other departures follow Table IV: any other departure (retirement, death, a move to another public plan, or a job at a private firm that is not an asset manager), retirement, and other jobs (another public plan or a non-asset-manager firm). "+CTXT+BTXT+SETXT,13,[34]+[11]*12,140)
r=4; c=2
for y,_ in YS:
    ws.merge_cells(start_row=r,start_column=c,end_row=r,end_column=c+5); x=ws.cell(row=r,column=c,value=YL[y]); x.font=F; x.alignment=CENTER
    for k in range(6): ws.cell(row=r,column=c+k).border=Border(top=thin)
    for fe in ['fy','ppd_id + fy']:
        ws.merge_cells(start_row=r+1,start_column=c,end_row=r+1,end_column=c+2); x=ws.cell(row=r+1,column=c,value=FL[fe]); x.font=F; x.alignment=CENTER
        for k,h in enumerate(['Conflicted x Shock','Other x Shock','Difference']): x=ws.cell(row=r+2,column=c+k,value=h); x.font=F; x.alignment=CENTER; x.border=Border(bottom=thin)
        c+=3
ws.cell(row=r,column=1).border=Border(top=thin); ws.cell(row=r+2,column=1).border=Border(bottom=thin); ws.row_dimensions[r+2].height=30
r+=3
for g,gl in GRP:
    gl={'oth_any':'Any other departure','oth_ret':'Retirement','oth_job':'Other job'}[g]
    ws.cell(row=r,column=1,value=gl).font=F
    for j,(y,fe,t) in enumerate(C6): x=R[('grp',g,y,fe)]; put(ws,r,2+j,reg('Placebo groups',y,FL[fe]+', '+gl,t,x[t],x['G']),FI if t=='diff' else F)
    r+=2
r=stats_row(ws,r,'Observations',[R[('grp','oth_any',y,fe)]['N'] for y,fe,t in C6],'#,##0'); bottom(ws,r,13)
r+=1; ws.cell(row=r,column=1,value="Coefficients and observations will match reghdfe; t-statistics and p-values can differ somewhat. Use the attached Stata code for the exact figures.").font=FI
wb.move_sheet('Estimates', offset=len(wb.sheetnames)-1)
out='/home/user/PensionData/referee_response/time_to_departure/shock_design_tests.xlsx'; wb.save(out); print('saved',out)
