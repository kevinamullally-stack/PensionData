import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter as L
D=pd.read_pickle('shock_es_fye.pkl'); RES=D['RES']; YS=D['YS']; FES2=D['FES2']; SINGLE=D['SINGLE']; BINP=D['BINP']
TNR='Times New Roman'; SZ=10.5
F=Font(name=TNR,size=SZ); FB=Font(name=TNR,size=SZ,bold=True); FI=Font(name=TNR,size=SZ,italic=True); BLUE=Font(name=TNR,size=SZ,color='0000FF')
thin=Side(style='thin'); CENTER=Alignment(horizontal='center',vertical='center',wrap_text=True); RIGHT=Alignment(horizontal='right'); JUST=Alignment(horizontal='justify',wrap_text=True,vertical='top')
SHADE=PatternFill('solid',fgColor='F1F0EC')
YL={'peer_adj_ret100':'Peer-Adjusted Return','b_AB':'Benchmark-Adjusted Return'}; FL={'fye_id':'Fiscal-Year-End FE','ppd_id + fye_id':'Plan and Fiscal-Year-End FE'}
COLS=[(y,fe) for y,_ in YS for fe,_ in FES2]
wb=Workbook(); E=wb.active; E.title='Estimates'
E['A1']='Estimates behind the tables. Blue = regression output (coefficient, SE, minimum cluster count). t and p are formulas; p uses the t distribution with (min clusters - 1) df. SE double clustered by plan and year, with the eigenvalue fix.'; E['A1'].font=FI
for j,h in enumerate(['Event study','Dependent variable','Specification','Term','Coefficient','SE','Clusters (min)','t','p'],1): E.cell(row=3,column=j,value=h).font=FB
for c,w in zip('ABCD',[22,26,28,44]): E.column_dimensions[c].width=w
er=[4]
def reg(es,y,fe,term,r,G):
    i=er[0]; er[0]+=1
    for j,v in enumerate([es,YL[y],FL[fe],term],1): E.cell(row=i,column=j,value=v).font=F
    E.cell(row=i,column=5,value=r['est']).font=BLUE; E.cell(row=i,column=6,value=r['se']).font=BLUE; E.cell(row=i,column=7,value=G).font=BLUE
    E.cell(row=i,column=8,value=f'=IF(F{i}>0,E{i}/F{i},"")').font=F; E.cell(row=i,column=9,value=f'=IF(H{i}="","",TDIST(ABS(H{i}),G{i}-1,2))').font=F
    for j in [5,6,8,9]: E.cell(row=i,column=j).number_format='0.0000'
    return dict(coef=f'Estimates!E{i}',t=f'Estimates!H{i}',p=f'Estimates!I{i}')
def cf(ref): return f'=IF({ref["p"]}="","",TEXT({ref["coef"]},"0.000")&IF({ref["p"]}<0.01,"***",IF({ref["p"]}<0.05,"**",IF({ref["p"]}<0.1,"*",""))))'
def tf(ref): return f'=IF({ref["t"]}="","","("&TEXT({ref["t"]},"0.00")&")")'
def put(ws,r,c,ref,fnt=F):
    a=ws.cell(row=r,column=c,value=cf(ref)); a.font=fnt; a.alignment=RIGHT; b=ws.cell(row=r+1,column=c,value=tf(ref)); b.font=fnt; b.alignment=RIGHT
def pcell(ws,r,c,ref,fnt=F):
    x=ws.cell(row=r,column=c,value=f'={ref["p"]}'); x.font=fnt; x.alignment=RIGHT; x.number_format='0.000'
def build(es,name,title,periods,labfn,note):
    ws=wb.create_sheet(name); ws.sheet_view.showGridLines=False; ws.column_dimensions['A'].width=50
    for c in range(2,6): ws.column_dimensions[L(c)].width=16
    ws['A1']=title; ws['A1'].font=FB
    ws.merge_cells('A2:E2'); ws['A2']=note; ws['A2'].font=F; ws['A2'].alignment=JUST; ws.row_dimensions[2].height=14*(len(note)//95+2)
    r=4; c=2
    for y,_ in YS:
        ws.merge_cells(start_row=r,start_column=c,end_row=r,end_column=c+1); x=ws.cell(row=r,column=c,value=YL[y]); x.font=F; x.alignment=CENTER
        for k in range(2): ws.cell(row=r,column=c+k).border=Border(top=thin)
        c+=2
    ws.cell(row=r,column=1).border=Border(top=thin); r+=1
    for j,(y,fe) in enumerate(COLS): x=ws.cell(row=r,column=2+j,value=FL[fe]); x.font=F; x.alignment=CENTER; x.border=Border(bottom=thin)
    ws.cell(row=r,column=1).border=Border(bottom=thin); ws.row_dimensions[r].height=30; r+=1
    # test block first
    ws.cell(row=r,column=1,value='Panel A. Shock-period years vs. non-shock years').font=FB; r+=1
    rows=[('Average effect of conflict, shock-period years (FY2009-2013)','lvl_shock',F,False),
          ('Average effect of conflict, non-shock years','lvl_non',F,False),
          ('Difference: shock-period minus non-shock years','d_non',FB,True)]
    for lab,key,fnt,withp in rows:
        ws.cell(row=r,column=1,value=lab).font=fnt
        refs=[]
        for j,(y,fe) in enumerate(COLS): x=RES[(es,y,fe)]; ref=reg(name,y,fe,lab,x[key],x['G']); put(ws,r,2+j,ref,fnt); refs.append(ref)
        r+=2
        if withp:
            ws.cell(row=r,column=1,value='p-value of the difference').font=fnt
            for j,ref in enumerate(refs): pcell(ws,r,2+j,ref,fnt)
            r+=1
    r+=1; ws.cell(row=r,column=1,value='Panel B. Non-shock years split into pre-shock and post-shock').font=FB; r+=1
    rows=[('Average effect of conflict, pre-shock years (FY2002-2008)','lvl_pre',F,False),
          ('Average effect of conflict, post-shock years (FY2014-2020)','lvl_post',F,False),
          ('Difference: shock-period minus pre-shock years','d_pre',FI,True),
          ('Difference: shock-period minus post-shock years','d_post',FI,True)]
    for lab,key,fnt,withp in rows:
        ws.cell(row=r,column=1,value=lab).font=fnt
        refs=[]
        for j,(y,fe) in enumerate(COLS): x=RES[(es,y,fe)]; ref=reg(name,y,fe,lab,x[key],x['G']); put(ws,r,2+j,ref,fnt); refs.append(ref)
        r+=2
        if withp:
            ws.cell(row=r,column=1,value='p-value of the difference').font=fnt
            for j,ref in enumerate(refs): pcell(ws,r,2+j,ref,fnt)
            r+=1
    r+=1; ws.cell(row=r,column=1,value='Panel C. Event-study coefficients').font=FB; r+=1
    ws.cell(row=r,column=1,value='Conflicted CIO (effect in the base period)').font=F
    for j,(y,fe) in enumerate(COLS): x=RES[(es,y,fe)]; put(ws,r,2+j,reg(name,y,fe,'Conflicted CIO',x['conf'],x['G']))
    r+=2
    for nm,yrs,grp in periods:
        if nm is None: continue
        ws.cell(row=r,column=1,value=labfn(nm)).font=F
        for j,(y,fe) in enumerate(COLS): x=RES[(es,y,fe)]; put(ws,r,2+j,reg(name,y,fe,labfn(nm),x[nm],x['G']))
        if grp=='shock':
            for rr in (r,r+1):
                for cc in range(1,6): ws.cell(row=rr,column=cc).fill=SHADE
        r+=2
    r+=1
    ws.cell(row=r,column=1,value='Observations').font=F
    for j,(y,fe) in enumerate(COLS): x=ws.cell(row=r,column=2+j,value=RES[(es,y,fe)]['N']); x.font=F; x.alignment=RIGHT; x.number_format='#,##0'
    r+=1
    for c in range(1,6): ws.cell(row=r-1,column=c).border=Border(bottom=thin)
    ws.page_setup.orientation='portrait'; ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
COMMON=("Each column is one regression of performance on Conflicted CIO and Conflicted CIO interacted with period indicators (the base period omitted), the Table VIII controls, and the fixed effects shown; standard errors are double clustered by plan and year. "
        "The average effect of conflict in a group of years is Conflicted CIO plus the average of that group's interaction coefficients, weighting each coefficient by the number of fiscal years it covers in the FY2002-2020 sample and counting the base period as zero; the differences therefore do not depend on which period is omitted. "
        "Shock-period years are FY2009-2013 (Shock Period = 1); non-shock years are FY2002-2008 and FY2014-2020. Benchmark-Adjusted Return is the plan's return minus the updated policy benchmark (reported, or inferred and graded A or B). "
        "Coefficients marked with ***, **, and * are significant at the 1%, 5%, and 10% level, respectively.")
def lab1(n): return {'es_pre':'Conflicted CIO x FY2002-2004','es_post':'Conflicted CIO x FY2018-2020'}.get(n,'Conflicted CIO x FY'+n[3:])
BL={'b0102':'2002 (2001-02 bin)','b0304':'2003-04','b0506':'2005-06','b0910':'2009-10','b1113':'2011-13','b1415':'2014-15','b1617':'2016-17','b1820':'2018-20'}
build('single','Single years','Shock-Period vs. Non-Shock Coefficients: Event Study With FY2008 Base',SINGLE,lab1,
      "The event study has an indicator for each fiscal year 2005-2017 (FY2008 omitted) and endpoint bins for FY2002-2004 and FY2018-2020. "+COMMON)
build('bins','Two-year bins','Shock-Period vs. Non-Shock Coefficients: Event Study With Two-Year Bins (FY2007-08 Base)',BINP,lambda n:'Conflicted CIO x FY'+BL[n],
      "The event study uses two-year periods (FY2007-08 omitted; the shock period split into FY2009-10 and FY2011-13; FY2018-20 pooled; the 2001-02 bin contains only FY2002 in the regression sample). "+COMMON)
wb.move_sheet('Estimates', offset=2)
out='/home/user/PensionData/referee_response/time_to_departure/shock_eventstudy_tests_fye1.xlsx'; wb.save(out); print('saved',out)
