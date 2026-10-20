import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
D=pd.read_pickle('returns_spell_slope.pkl'); A=D['A']; X=D['X']; BY=D['BYLEN']; C=D['C']; RATE=D['RATE']; LEV=D['LEV']; YS=D['YS']
TNR='Times New Roman'; SZ=10.5
F=Font(name=TNR,size=SZ); FB=Font(name=TNR,size=SZ,bold=True); FI=Font(name=TNR,size=SZ,italic=True)
thin=Side(style='thin'); CENTER=Alignment(horizontal='center',vertical='center',wrap_text=True); RIGHT=Alignment(horizontal='right',vertical='center'); LEFT=Alignment(horizontal='left',vertical='center',wrap_text=True); JUST=Alignment(horizontal='justify',wrap_text=True,vertical='top')
def st(p): return '' if p is None or np.isnan(p) else '***' if p<0.01 else '**' if p<0.05 else '*' if p<0.1 else ''
def cp(v,p,dec=3): return f"{v:+.{dec}f}{st(p)} ({p:.3f})"
W=[50,18,18,18,14,18,9,9,18,13,12]; NC=len(W)
def header(ws,r,H,h=58):
    for j,x_ in enumerate(H): x=ws.cell(row=r,column=1+j,value=x_); x.font=F; x.alignment=CENTER; x.border=Border(top=thin,bottom=thin)
    ws.row_dimensions[r].height=h
def rowvals(ws,r,lab,vals,font=None):
    ws.cell(row=r,column=1,value=lab).font=font or (F if lab.startswith('  ') else FB); ws.cell(row=r,column=1).alignment=LEFT
    for j,v in enumerate(vals): x=ws.cell(row=r,column=2+j,value=v); x.font=font or F; x.alignment=RIGHT
def rule(ws,r):
    for c in range(1,NC+1): ws.cell(row=r,column=c).border=Border(bottom=thin)
wb=Workbook(); first=True
for y,yl in YS:
    ws=wb.active if first else wb.create_sheet(); first=False
    ws.title=yl.replace('Return','return'); ws.sheet_view.showGridLines=False
    for j,w in enumerate(W): ws.column_dimensions[L(1+j)].width=w
    ws['A1']=f'Spell-Level Slope of the {yl}: Conflicted vs. Other Departing CIOs'; ws['A1'].font=FB
    note=(f"Sample: CIOs who have left (curr = 0) in the Table III sample ({D['nP'][y]:,} plan-years with {yl.lower()}s, {D['nsp'][y]} CIO-plan spells). "
          f"The spell-level slope is each CIO-plan spell's OLS slope of the {yl.lower()} (%) on fiscal year, in percentage points per year, winsorized at 1% and 99%; it is the same construction as the Fee Slope. Every spell counts once. "
          "Panel A compares mean slopes of conflicted and other departing CIOs: p-values test that each group's mean is zero (signed-rank p-value for the conflicted median) and that the groups differ (t-test, t-test allowing unequal variances, rank-sum test); 'weighted' weights each spell by its plan-years (standard errors clustered by plan); 'share rising' is the share of spells with a non-negative slope. "
          "Rows below the three-year row use spells with at least three years: the return net of the mean of all Table III plan-years with the same fiscal-year-end date; one observation per CIO (returns averaged across plans within a CIO-year); slopes not winsorized; and dropping the conflicted spells with the most negative or positive slopes. "
          "Panel B regresses the slope on Conflicted CIO and the spell's first-year return (mean reversion) and/or the spell's own slopes of the equity, fixed income and alternatives shares (equal weights, standard errors clustered by plan). Panel C reports the first-year, mean and last-year return of the same spells. Panel D reports mean slopes by spell length. "
          "Panel E reports designs that weight by plan-year (CIO-spell fixed effects, standard errors double clustered by plan and year) and the first-to-second-half change divided by the years between the midpoints of the two halves. ***, **, and * denote significance at the 1%, 5%, and 10% level.")
    ws.merge_cells(f'A2:{L(NC)}2'); ws['A2']=note; ws['A2'].font=F; ws['A2'].alignment=JUST; ws.row_dimensions[2].height=14*(len(note)//(sum(W)*1.3)+1)
    r=4; ws.cell(row=r,column=1,value='Panel A. Spell-level slope (pp per year)').font=FB; r+=1
    header(ws,r,['Spells included','Conflicted: mean slope (p vs. 0)','Conflicted: median (signed-rank p)','Other departing CIOs: mean slope (p vs. 0)','Other: median','Conflicted minus other (t-test p)','Welch p','Rank-sum p','Weighted by plan-years (p)','Share rising (conflicted / other)','Spells (conflicted / other)']); r+=1
    for k,lab in [('min2','At least 2 years'),('min3','At least 3 years'),('min4','At least 4 years'),('min6','At least 6 years'),('min8','At least 8 years'),
                  ('dm','  Net of fiscal-year-end mean'),('cio','  One observation per CIO'),('nowins','  Slopes not winsorized'),
                  ('dropneg3','  Drop the 3 most negative conflicted slopes'),('droppos3','  Drop the 3 most positive conflicted slopes'),('dropneg5','  Drop the 5 most negative conflicted slopes'),('droppos5','  Drop the 5 most positive conflicted slopes')]:
        q=A[(y,k)]
        rowvals(ws,r,lab,[cp(q['mc'],q['p1c']),cp(q['medc'],q['psr']),cp(q['mo'],q['p1o']),f"{q['medo']:+.3f}",cp(q['diff'],q['p']),f"{q['pw']:.3f}",f"{q['prs']:.3f}",cp(q['wdiff'],q['wp']),f"{1-q['negc']:.2f} / {1-q['nego']:.2f}",f"{q['nc']} / {q['no']}"],
                font=FB if k=='min3' else None); r+=1
    rule(ws,r-1); r+=1
    ws.cell(row=r,column=1,value='Panel B. Slope regressed on Conflicted CIO and controls (at least 3 years)').font=FB; r+=1
    header(ws,r,['Controls','Conflicted CIO (p)','First-year return (p)','Equity share slope (p)','Fixed income share slope (p)','Alternatives share slope (p)','Spells','','','',''],44); r+=1
    for k,lab in [('none','None'),('y0','First-year return'),('none_a','None (spells with allocation data)'),('alloc','Own allocation slopes'),('both','First-year return and own allocation slopes')]:
        q=X[(y,k)]; cf=q['coefs']; g=lambda v: cp(*cf[v]) if v in cf else ''
        vals=[cp(q['diff'],q['p']),g('y0'),g('s_eq_pct'),g('s_fi_pct'),g('s_alt_pct'),f"{q['n']}"]
        rowvals(ws,r,'  '+lab,vals); r+=1
    rule(ws,r-1); r+=1
    ws.cell(row=r,column=1,value='Panel C. Return levels of the same spells (at least 3 years)').font=FB; r+=1
    header(ws,r,['','Conflicted mean','Other mean','Conflicted minus other (p)','Spells (conflicted / other)','','','','','',''],30); r+=1
    for k,lab in [('first','First year of the spell'),('mean','Spell mean'),('last','Last year of the spell')]:
        q=LEV[y][k]; rowvals(ws,r,'  '+lab,[f"{q['c']:+.3f}",f"{q['o']:+.3f}",cp(q['c']-q['o'],q['p']),f"{q['nc']} / {q['no']}"]); r+=1
    rule(ws,r-1); r+=1
    ws.cell(row=r,column=1,value='Panel D. Mean slope by spell length (at least 3 years)').font=FB; r+=1
    header(ws,r,['Plan-years in the spell','Conflicted mean slope','Conflicted spells','Other mean slope','Other spells','','','','','',''],30); r+=1
    B=BY[y]
    for lv in B.index:
        rowvals(ws,r,'  '+str(lv),[f"{B.loc[lv,('mean',1.0)]:+.3f}",f"{int(B.loc[lv,('count',1.0)])}",f"{B.loc[lv,('mean',0.0)]:+.3f}",f"{int(B.loc[lv,('count',0.0)])}"]); r+=1
    rule(ws,r-1); r+=1
    ws.cell(row=r,column=1,value='Panel E. Designs weighted by plan-year, and the per-year halves change').font=FB; r+=1
    header(ws,r,['Design','Conflicted extra change','t-statistic','p-value','Observations','Other departing CIOs (p)','','','','',''],44); r+=1
    for k,lab in [('pool','Pooled trend: Conflicted x fiscal year (pp per year), spell FE'),('pool_fye','Pooled trend: Conflicted x fiscal year (pp per year), spell and FYE FE'),
                  ('halves','Second minus first half, conflicted minus other, spell FE'),('halves_fye','Second minus first half, conflicted minus other, spell and FYE FE')]:
        q=C[(y,k)]; b=q['b']; oth=cp(q['oth']['est'],q['oth']['p']) if 'oth' in q else ''
        rowvals(ws,r,'  '+lab,[f"{b['est']:+.3f}{st(b['p'])}",f"{b['t']:.2f}",f"{b['p']:.3f}",f"{q['N']:,}",oth]); ws.row_dimensions[r].height=28; r+=1
    q=RATE[y]
    rowvals(ws,r,f"  Second-minus-first change per year between half midpoints, equal weight per spell (years apart: {q['gapc']:.2f} / {q['gapo']:.2f})",
            [cp(q['diff'],q['p']),'',f"{q['p']:.3f}",f"{q['nc']} / {q['no']} spells",cp(q['mo'],q['p1o'])]); ws.row_dimensions[r].height=28; r+=1
    rule(ws,r-1)
    ws.page_setup.orientation='landscape'; ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
out='/home/user/PensionData/referee_response/time_to_departure/returns_spell_slope.xlsx'; wb.save(out); print('saved',out)
