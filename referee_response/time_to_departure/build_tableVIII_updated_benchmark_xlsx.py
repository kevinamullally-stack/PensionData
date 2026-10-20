import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
D=pd.read_pickle('tableVIII_bench.pkl'); T8=D['T8']; YS=D['YS']; FES=D['FES']; PLAN=D['PLAN']
TNR='Times New Roman'; SZ=10.5
F=Font(name=TNR,size=SZ); FB=Font(name=TNR,size=SZ,bold=True); FI=Font(name=TNR,size=SZ,italic=True); BLUE=Font(name=TNR,size=SZ,color='0000FF')
thin=Side(style='thin'); CENTER=Alignment(horizontal='center',vertical='center',wrap_text=True); RIGHT=Alignment(horizontal='right'); JUST=Alignment(horizontal='justify',wrap_text=True,vertical='top')
YL={'peer_adj_ret100':'Peer-adj. Ret (t)','b_AB':'Bench-adj. Ret (t)'}
wb=Workbook(); E=wb.active; E.title='Estimates'
E['A1']='Estimates behind the table. Blue = regression output (coefficient, SE, minimum cluster count). t and p are formulas; p uses the t distribution with (clusters - 1) df.'; E['A1'].font=FI
for j,h in enumerate(['Panel','Dependent variable','Column','Term','Coefficient','SE','Clusters (min)','t','p'],1): E.cell(row=3,column=j,value=h).font=FB
for c,w in zip('ABCD',[10,26,10,40]): E.column_dimensions[c].width=w
er=[4]
def reg(pan,y,col,term,r,G):
    i=er[0]; er[0]+=1
    for j,v in enumerate([pan,y,col,term],1): E.cell(row=i,column=j,value=v).font=F
    E.cell(row=i,column=5,value=r['est']).font=BLUE; E.cell(row=i,column=6,value=r['se']).font=BLUE; E.cell(row=i,column=7,value=G).font=BLUE
    E.cell(row=i,column=8,value=f'=IF(F{i}>0,E{i}/F{i},"")').font=F; E.cell(row=i,column=9,value=f'=IF(H{i}="","",TDIST(ABS(H{i}),G{i}-1,2))').font=F
    for j in [5,6,8,9]: E.cell(row=i,column=j).number_format='0.0000'
    return dict(coef=f'Estimates!E{i}',t=f'Estimates!H{i}',p=f'Estimates!I{i}')
def cf(ref): return f'=IF({ref["p"]}="","",TEXT({ref["coef"]},"0.000")&IF({ref["p"]}<0.01,"***",IF({ref["p"]}<0.05,"**",IF({ref["p"]}<0.1,"*",""))))'
def tf(ref): return f'=IF({ref["t"]}="","","("&TEXT({ref["t"]},"0.00")&")")'
ws=wb.create_sheet('Table VIII',0); ws.sheet_view.showGridLines=False; ws.column_dimensions['A'].width=50
for c in range(2,7): ws.column_dimensions[L(c)].width=14
ws['A1']='Table VIII: Exogenous Shock - Conflicted CIOs During the Loglisci Case, Peer- and Benchmark-Adjusted Returns'; ws['A1'].font=FB
note=("This table reports results of tests that use a natural shock to CIOs' susceptibility of being conflicted. Specifically, we compare the impact of Conflicted CIO during the period of March 2009 to March 2014 when the case surrounding former NY Common Fund CIO David Loglisci was still unsettled. "
      "Columns (1)-(3) are the published specifications; columns (4) and (5) replace year fixed effects with fiscal-year-end date fixed effects. Panel A uses Peer-Adjusted Return; its columns (1)-(3) reproduce the published Table VIII. Panel B uses Benchmark-Adjusted Return computed as the plan's return minus the updated policy benchmark (the benchmark return reported in the plan's financial report or, where none is stated, an inferred value graded A or B). "
      "Shock Period is an indicator variable equal to 1 if a plan fiscal year ends between March 2009 and March 2014 and 0 otherwise; it is a function of the fiscal-year-end date, so it is absorbed in columns (4) and (5), and in the benchmark-adjusted sample it is also absorbed by the year fixed effects in columns (1) and (3). The other independent variables are as defined in Table I. "
      "The standard errors are double clustered by plan and year. Coefficients marked with ***, **, and * are significant at the 1%, 5%, and 10% level, respectively. "
      "The last rows of each panel report t-statistics and the F-test p-value with standard errors clustered by plan only: in Panel B, column (1), the two-way clustered variance of Conflicted + Interaction is not positive definite, and the standard adjustment for that case yields an artificially small standard error (p = 0.006); clustered by plan, p = 0.41.")
ws.merge_cells('A2:F2'); ws['A2']=note; ws['A2'].font=F; ws['A2'].alignment=JUST; ws.row_dimensions[2].height=14*(len(note)//115+2)
r=4
for c in range(6): x=ws.cell(row=r,column=1+c,value=['','(1)','(2)','(3)','(4)','(5)'][c]); x.font=F; x.alignment=CENTER; x.border=Border(top=thin)
r+=1
for pan,(y,_) in zip(['A','B'],YS):
    ws.cell(row=r,column=1,value=f"Panel {pan}. {'Peer-Adjusted Return' if pan=='A' else 'Benchmark-Adjusted Return (updated benchmark)'}").font=FB
    ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=6); x=ws.cell(row=r,column=2,value=YL[y]); x.font=F; x.alignment=CENTER
    for c in range(2,7): ws.cell(row=r,column=c).border=Border(bottom=thin)
    r+=1
    for lab,k in [('Conflicted CIO (0/1)','conf'),('Conflicted x Shock Period (0/1)','cxs'),('Shock Period (0/1)','shock')]:
        ws.cell(row=r,column=1,value=lab).font=F
        for j,(fe,*_) in enumerate(FES):
            x=T8[(y,fe)][k]
            if np.isnan(x['est']):
                a=ws.cell(row=r,column=2+j,value='absorbed'); a.font=FI; a.alignment=RIGHT; continue
            ref=reg(pan,y,f'({j+1})',lab,x,T8[(y,fe)]['G'])
            a=ws.cell(row=r,column=2+j,value=cf(ref)); a.font=F; a.alignment=RIGHT; b=ws.cell(row=r+1,column=2+j,value=tf(ref)); b.font=F; b.alignment=RIGHT
        r+=2
    r+=1
    for lab,vals,fmt in [('Controls',['YES']*5,'@'),('Plan FE',[f[1] for f in FES],'@'),('Year FE',[f[2] for f in FES],'@'),('Fiscal-Year-End FE',[f[3] for f in FES],'@'),
                         ('Observations',[T8[(y,f[0])]['N'] for f in FES],'#,##0'),('Adj. R-squared',[T8[(y,f[0])]['adjr2'] for f in FES],'0.000')]:
        ws.cell(row=r,column=1,value=lab).font=F
        for j,v in enumerate(vals): x=ws.cell(row=r,column=2+j,value=v.upper() if isinstance(v,str) else v); x.font=F; x.alignment=CENTER if isinstance(v,str) else RIGHT; x.number_format=fmt
        r+=1
    ws.cell(row=r,column=1,value='Conflicted + Interaction').font=F
    refs=[]
    for j,(fe,*_) in enumerate(FES):
        x=T8[(y,fe)]; ref=reg(pan,y,f'({j+1})','Conflicted + Interaction',x['sum'],x['G']); refs.append(ref)
        c=ws.cell(row=r,column=2+j,value=f'={ref["coef"]}'); c.font=F; c.alignment=RIGHT; c.number_format='0.000'
    r+=1
    ws.cell(row=r,column=1,value='F-Test of Sum ≠ 0 (p-value)').font=F
    for j,ref in enumerate(refs): c=ws.cell(row=r,column=2+j,value=f'={ref["p"]}'); c.font=F; c.alignment=RIGHT; c.number_format='0.000'
    r+=1
    for lab,k in [('t, Conflicted CIO (clustered by plan)','conf'),('t, Conflicted x Shock Period (clustered by plan)','cxs'),('p, F-Test of Sum ≠ 0 (clustered by plan)','sum')]:
        ws.cell(row=r,column=1,value=lab).font=FI
        for j,(fe,*_) in enumerate(FES):
            x=PLAN[(y,fe)][k]; ref=reg(pan+' (plan-clustered)',y,f'({j+1})',lab,x,PLAN[(y,fe)]['G'])
            c=ws.cell(row=r,column=2+j,value=f'={ref["p"] if k=="sum" else ref["t"]}'); c.font=FI; c.alignment=RIGHT; c.number_format='0.000' if k=='sum' else '0.00'
        r+=1
    for c in range(1,7): ws.cell(row=r-1,column=c).border=Border(bottom=thin)
    r+=1
ws.cell(row=r,column=1,value="Coefficients and observations match reghdfe; two-way clustered t-statistics and p-values can differ slightly (Panel A reproduces the published t-statistics to within 0.05). Use the attached Stata code for the exact figures.").font=FI
ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=6); ws.cell(row=r,column=1).alignment=JUST; ws.row_dimensions[r].height=30
ws.page_setup.orientation='portrait'; ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
out='/home/user/PensionData/referee_response/time_to_departure/tableVIII_updated_benchmark.xlsx'; wb.save(out); print('saved',out)
