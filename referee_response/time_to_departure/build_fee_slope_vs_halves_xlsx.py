import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
D=pd.read_pickle('fee_reconcile.pkl'); O=D['OUT']; BL=D['BYLEN']; IM=D['IMPL']; MED=D['MED']; HV=D['HV']
TNR='Times New Roman'; SZ=10.5
F=Font(name=TNR,size=SZ); FB=Font(name=TNR,size=SZ,bold=True); FI=Font(name=TNR,size=SZ,italic=True); BLUE=Font(name=TNR,size=SZ,color='0000FF')
thin=Side(style='thin'); CENTER=Alignment(horizontal='center',vertical='center',wrap_text=True); RIGHT=Alignment(horizontal='right'); LEFT=Alignment(horizontal='left',vertical='center',wrap_text=True); JUST=Alignment(horizontal='justify',wrap_text=True,vertical='top')
def st(p): return '' if p is None or np.isnan(p) else '***' if p<0.01 else '**' if p<0.05 else '*' if p<0.1 else ''
wb=Workbook(); ws=wb.active; ws.title='Reconciliation'; ws.sheet_view.showGridLines=False
for j,w in enumerate([52,13,13,15,10,15,10,12]): ws.column_dimensions[L(1+j)].width=w
ws['A1']='Fee Ratio: Spell-Level Fee Slope vs. First-Half/Second-Half Comparisons'; ws['A1'].font=FB
note=(f"Sample: CIOs who have left (curr = 0) in the Table III sample with Fee Ratio data ({D['nD']:,} plan-years, {D['nsp']} CIO-plan spells). Fee Ratio is total external fees divided by plan assets, in percent. "
      "Panel A follows the Table IX construction without allocation controls: each spell's OLS slope of Fee Ratio on fiscal year (percentage points per year), winsorized at 1% and 99%, compared between conflicted and other CIOs with a t-test (equal weight per spell) and with a regression weighting each spell by its number of plan-years (standard errors clustered by plan). "
      "Panel B reports mean slopes by spell length. Panel C reports designs that weight by plan-year: a pooled regression of Fee Ratio on Conflicted CIO x fiscal year with CIO-spell fixed effects (and fiscal-year-end fixed effects), the change from the first to the second half of tenure within spell for conflicted relative to other departing CIOs, and the halves comparison against all other plan-years from halves_v3_fye1.xlsx (Panel A, Fee Ratio). "
      "Panel C regressions have standard errors double clustered by plan and year. ***, **, and * denote significance at the 1%, 5%, and 10% level.")
ws.merge_cells('A2:H2'); ws['A2']=note; ws['A2'].font=F; ws['A2'].alignment=JUST; ws.row_dimensions[2].height=14*(len(note)//120+2)
def hdr(r,cols):
    for j,h in enumerate(cols): x=ws.cell(row=r,column=1+j,value=h); x.font=F; x.alignment=CENTER; x.border=Border(top=thin,bottom=thin)
    ws.row_dimensions[r].height=44
def val(r,c,v,fmt='0.0000',fnt=F,stars=None):
    if stars is not None:
        x=ws.cell(row=r,column=c,value=f"{v:+.4f}{stars}"); x.font=fnt; x.alignment=RIGHT; return
    x=ws.cell(row=r,column=c,value=v); x.font=fnt; x.alignment=RIGHT; x.number_format=fmt
r=4; ws.cell(row=r,column=1,value='Panel A. Spell-level Fee Slope (no allocation controls)').font=FB; r+=1
hdr(r,['Spells included','Conflicted mean slope','Other mean slope','Difference (equal weight)','p-value','Difference (weighted by plan-years)','p-value','Spells (conflicted / other)']); r+=1
for k,lab in [('base2','At least 2 years of fee data'),('base3','At least 3 years'),('dm3','At least 3 years, Fee Ratio net of the fiscal-year mean'),('min4','At least 4 years'),('min6','At least 6 years'),('min8','At least 8 years')]:
    a=O[('slope',k)]; w=O[('slope_w',k)]
    ws.cell(row=r,column=1,value=lab).font=F
    val(r,2,a['conf']); val(r,3,a['oth']); val(r,4,a['diff'],stars=st(a['p']),fnt=FB); val(r,5,a['p'],'0.000',BLUE); val(r,6,w['diff'],stars=st(w['p']),fnt=FB); val(r,7,w['p'],'0.000',BLUE)
    x=ws.cell(row=r,column=8,value=f"{a['nc']} / {a['no']}"); x.font=F; x.alignment=RIGHT; r+=1
for mo,lab in [(3,'Medians, at least 3 years (rank-sum test)'),(6,'Medians, at least 6 years (rank-sum test)')]:
    m=MED[mo]; ws.cell(row=r,column=1,value=lab).font=FI
    val(r,2,m['mc'],fnt=FI); val(r,3,m['mo_'],fnt=FI); val(r,4,m['mc']-m['mo_'],stars=st(m['p']),fnt=FI); val(r,5,m['p'],'0.000',BLUE)
    x=ws.cell(row=r,column=6,value=f"share rising: {m['pc']:.2f} vs {m['po']:.2f}"); x.font=FI; x.alignment=RIGHT; ws.merge_cells(start_row=r,start_column=6,end_row=r,end_column=8); r+=1
for c in range(1,9): ws.cell(row=r-1,column=c).border=Border(bottom=thin)
r+=1; ws.cell(row=r,column=1,value='Panel B. Mean Fee Slope by spell length (at least 3 years)').font=FB; r+=1
hdr(r,['Plan-years of fee data in the spell','Conflicted mean slope','Conflicted spells','Other mean slope','Other spells','','','']); r+=1
for lenlab in BL.index:
    ws.cell(row=r,column=1,value=str(lenlab)).font=F
    val(r,2,float(BL.loc[lenlab,('mean',1.0)])); val(r,3,int(BL.loc[lenlab,('count',1.0)]),'0'); val(r,4,float(BL.loc[lenlab,('mean',0.0)])); val(r,5,int(BL.loc[lenlab,('count',0.0)]),'0'); r+=1
for c in range(1,9): ws.cell(row=r-1,column=c).border=Border(bottom=thin)
r+=1; ws.cell(row=r,column=1,value='Panel C. Designs weighted by plan-year').font=FB; r+=1
hdr(r,['Design','Conflicted extra change','t-statistic','p-value','Observations','','','']); r+=1
for k,lab in [('pool_noyear','Pooled trend: Conflicted x fiscal year (pp per year), spell FE'),('pool_year','Pooled trend: Conflicted x fiscal year (pp per year), spell and FYE FE'),
              ('halves_noyear','Second minus first half, conflicted minus other departing CIOs, spell FE'),('halves_year','Second minus first half, conflicted minus other departing CIOs, spell and FYE FE')]:
    x=O[k]['b']; ws.cell(row=r,column=1,value=lab).font=F; ws.cell(row=r,column=1).alignment=LEFT; ws.row_dimensions[r].height=28
    val(r,2,x['est'],stars=st(x['p']),fnt=FB); val(r,3,x['t'],'0.00'); val(r,4,x['p'],'0.000',BLUE); val(r,5,O[k]['N'],'#,##0'); r+=1
for k,lab in [('fye','Second minus first half vs. all other plan-years, FYE FE (halves_v3_fye1)'),('plan','Second minus first half vs. all other plan-years, plan and FYE FE (halves_v3_fye1)')]:
    x=HV[k]; ws.cell(row=r,column=1,value=lab).font=F; ws.cell(row=r,column=1).alignment=LEFT; ws.row_dimensions[r].height=28
    val(r,2,x['est'],stars=st(x['p']),fnt=FB); val(r,3,x['t'],'0.00'); val(r,4,x['p'],'0.000',BLUE); r+=1
for c in range(1,9): ws.cell(row=r-1,column=c).border=Border(bottom=thin)
r+=1
imp=O[('slope','base3')]['diff']*IM['gap']
txt=(f"Benchmark: the equal-weighted slope difference in Panel A (at least 3 years, {O[('slope','base3')]['diff']:+.4f} pp per year) times the average {IM['gap']:.2f}-year distance between the midpoints of conflicted CIOs' first and second halves implies a second-half difference of about {imp:+.3f} pp. "
     f"The plan-year-weighted designs in Panel C estimate {O['halves_year']['b']['est']:+.3f} (halves within spell, spell and FYE FE) and {O['pool_year']['b']['est']*IM['gap']:+.3f} (pooled trend times {IM['gap']:.2f} years).")
ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=8); ws.cell(row=r,column=1,value=txt).font=FI; ws.cell(row=r,column=1).alignment=JUST; ws.row_dimensions[r].height=44
ws.page_setup.orientation='landscape'; ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
out='/home/user/PensionData/referee_response/time_to_departure/fee_slope_vs_halves.xlsx'; wb.save(out); print('saved',out)
