import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
D=pd.read_pickle('halves_v3.pkl'); R=D['R']; PAIR=D['PAIR']; OUT=D['OUT']
GROUPS=[('Returns',['peer_adj_ret100','b_AB']),('Fees',['fee_ratio_pct','exp_ratio_pct']),('Fee rates by asset class',['eqfee100','fifee100','refee100','pefee100','hffee100','alt_fee_share']),('Allocations',['alt_pct','pe_pct','hf_pct'])]
LAB=dict(OUT)
TNR='Times New Roman'; SZ=10.5
F=Font(name=TNR,size=SZ); FB=Font(name=TNR,size=SZ,bold=True); FI=Font(name=TNR,size=SZ,italic=True); BLUE=Font(name=TNR,size=SZ,color='0000FF')
thin=Side(style='thin'); CENTER=Alignment(horizontal='center',vertical='center',wrap_text=True); RIGHT=Alignment(horizontal='right'); JUST=Alignment(horizontal='justify',wrap_text=True,vertical='top')
wb=Workbook(); E=wb.active; E.title='Estimates'
E['A1']='Estimates behind the tables. Blue = regression output (coefficient, SE, minimum cluster count). t and p are formulas; p uses the t distribution with (min clusters - 1) df. SE double clustered by plan and year, with the eigenvalue fix.'; E['A1'].font=FI
for j,h in enumerate(['Table','Outcome','Specification','Term','Coefficient','SE','Clusters (min)','t','p'],1): E.cell(row=3,column=j,value=h).font=FB
for c,w in zip('ABCD',[22,32,30,20]): E.column_dimensions[c].width=w
er=[4]
def reg(tab,y,spec,term,r,G):
    i=er[0]; er[0]+=1
    for j,v in enumerate([tab,LAB[y],spec,term],1): E.cell(row=i,column=j,value=v).font=F
    E.cell(row=i,column=5,value=None if np.isnan(r['est']) else r['est']).font=BLUE; E.cell(row=i,column=6,value=None if np.isnan(r['se']) else r['se']).font=BLUE; E.cell(row=i,column=7,value=G).font=BLUE
    E.cell(row=i,column=8,value=f'=IF(AND(ISNUMBER(E{i}),F{i}>0),E{i}/F{i},"")').font=F; E.cell(row=i,column=9,value=f'=IF(H{i}="","",TDIST(ABS(H{i}),G{i}-1,2))').font=F
    for j in [5,6,8,9]: E.cell(row=i,column=j).number_format='0.0000'
    return dict(coef=f'Estimates!E{i}',t=f'Estimates!H{i}',p=f'Estimates!I{i}')
def cf(ref): return f'=IF({ref["p"]}="","",TEXT({ref["coef"]},"0.000")&IF({ref["p"]}<0.01,"***",IF({ref["p"]}<0.05,"**",IF({ref["p"]}<0.1,"*",""))))'
def tf(ref): return f'=IF({ref["t"]}="","","("&TEXT({ref["t"]},"0.00")&")")'
def put(ws,r,c,ref,fnt=F):
    a=ws.cell(row=r,column=c,value=cf(ref)); a.font=fnt; a.alignment=RIGHT; b=ws.cell(row=r+1,column=c,value=tf(ref)); b.font=fnt; b.alignment=RIGHT
def pc(ws,r,c,ref,fnt=F): x=ws.cell(row=r,column=c,value=f'={ref["p"]}'); x.font=fnt; x.alignment=RIGHT; x.number_format='0.000'
def sheet(name,title,note,heads,widths,per=120):
    ws=wb.create_sheet(name); ws.sheet_view.showGridLines=False
    for j,w in enumerate(widths): ws.column_dimensions[L(1+j)].width=w
    ws['A1']=title; ws['A1'].font=FB; n=len(widths)
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=n); ws['A2']=note; ws['A2'].font=F; ws['A2'].alignment=JUST; ws.row_dimensions[2].height=14*(len(note)//per+2)
    r=4
    for lab,c0,c1 in heads[0]:
        if c1>c0: ws.merge_cells(start_row=r,start_column=c0,end_row=r,end_column=c1)
        x=ws.cell(row=r,column=c0,value=lab); x.font=F; x.alignment=CENTER
    for c in range(1,n+1): ws.cell(row=r,column=c).border=Border(top=thin)
    r+=1
    for j,h in enumerate(heads[1]):
        x=ws.cell(row=r,column=2+j,value=h); x.font=F; x.alignment=CENTER
    for c in range(1,n+1): ws.cell(row=r,column=c).border=Border(bottom=thin)
    ws.row_dimensions[r].height=44
    ws.page_setup.orientation='landscape'; ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    return ws, r+1
CTXT=("Controls for returns and fees: lagged return, Size, lagged funding ratio, lagged allocations, Financial Center, prior private experience, elite college, separate investment board and political appointees; allocation outcomes omit the lagged allocations. "
      "Lagged log CIO tenure and age are omitted from all specifications because tenure rises mechanically from the first to the second half. ")
DEF=("The second half of a CIO's tenure is the years in which tenure exceeds half of tenure at departure (tenure at departure = the CIO's maximum tenure at the plan plus any gap to the departure year). ")
SETXT="Standard errors are double clustered by plan and year. Coefficients marked with ***, **, and * are significant at the 1%, 5%, and 10% level, respectively."
# ---- Sheet A ----
ws,r=sheet('A. vs rest of sample','Conflicted CIOs, First vs. Second Half of Tenure, Relative to All Other Plan-Years',
 "Each row reports one regression on the Table III sample of the outcome on Conflicted CIO x First Half and Conflicted CIO x Second Half (non-conflicted plan-years are the omitted group; conflicted plan-years without departure information enter as a separate indicator), the controls, and fiscal-year-end date fixed effects (columns 1-4) or plan and fiscal-year-end date fixed effects (columns 5-6). "+DEF+CTXT+SETXT,
 [[('Fiscal-Year-End FE',2,5),('Plan and Fiscal-Year-End FE',6,7)],['Conflicted x First Half','Conflicted x Second Half','Second minus First','p-value','Second minus First','p-value','Observations']],[34,14,14,14,10,14,10,12])
for gl,ys in GROUPS:
    ws.cell(row=r,column=1,value=gl).font=FB; r+=1
    for y in ys:
        a=R[('split','fye',y)]; b=R[('split','plan',y)]
        ws.cell(row=r,column=1,value=LAB[y]).font=F
        put(ws,r,2,reg('A',y,'FYE FE','First half',a['first'],a['G'])); put(ws,r,3,reg('A',y,'FYE FE','Second half',a['second'],a['G']))
        ra=reg('A',y,'FYE FE','Second minus first',a['diff'],a['G']); put(ws,r,4,ra,FB); pc(ws,r,5,ra)
        rb=reg('A',y,'Plan + FYE FE','Second minus first',b['diff'],b['G']); put(ws,r,6,rb,FB); pc(ws,r,7,rb)
        x=ws.cell(row=r,column=8,value=a['N']); x.font=F; x.alignment=RIGHT; x.number_format='#,##0'
        r+=2
for c in range(1,9): ws.cell(row=r-1,column=c).border=Border(bottom=thin)
# ---- Sheet B ----
ws,r=sheet('B. within conflicted CIOs','Conflicted CIOs Compared With Themselves: Second Half of Tenure',
 f"Sample: conflicted CIOs who have left (curr = 0) with defined halves ({D['nC']} plan-years, {D['nCIO']} CIOs; {D['n12'][0]} first-half and {D['n12'][1]} second-half plan-years). Each cell is the coefficient on Second Half (0/1) from a regression of the outcome on Second Half, the controls, and the fixed effects shown; CIO fixed effects are CIO-by-plan spells, so columns 2-3 compare each CIO with themselves. Column 3 keeps CIOs with tenure at departure of at least four years ({D['nCIO4']} CIOs). "
 "Because a CIO's tenure and calendar time move together, CIO and fiscal-year-end fixed effects together leave only differences in when CIOs' halves fall to identify Second Half. "+DEF+CTXT+SETXT,
 [[('Second Half of Tenure (0/1)',2,4),('Observations',5,6)],['Fiscal-Year-End FE','CIO and Fiscal-Year-End FE','CIO and FYE FE, tenure >= 4','(1)','(2)']],[34,16,18,18,11,11])
for gl,ys in GROUPS:
    ws.cell(row=r,column=1,value=gl).font=FB; r+=1
    for y in ys:
        ws.cell(row=r,column=1,value=LAB[y]).font=F
        for j,(k,sl) in enumerate([('fye','FYE FE'),('cio','CIO + FYE FE'),('cio4','CIO + FYE FE, tenure>=4')]):
            x=R[('conf',k,y)]; put(ws,r,2+j,reg('B',y,sl,'Second half',x['second'],x['G']))
        for j,k in enumerate(['fye','cio']): x=ws.cell(row=r,column=5+j,value=R[('conf',k,y)]['N']); x.font=F; x.alignment=RIGHT; x.number_format='#,##0'
        r+=2
for c in range(1,7): ws.cell(row=r-1,column=c).border=Border(bottom=thin)
# ---- Sheet C ----
ws,r=sheet('C. CIO-level paired','CIO-Level Paired Comparison: Second-Half Mean minus First-Half Mean',
 "Each conflicted CIO who has left contributes one observation: the outcome averaged over plans within each CIO-year, then over the years in each half of tenure. Columns report the number of CIOs observed in both halves, the mean of the first- and second-half averages, the mean difference, p-values from a paired t-test and a Wilcoxon signed-rank test, and the share of CIOs whose second-half average exceeds their first-half average. "
 "These comparisons do not adjust for calendar time, so they include sample-wide trends (fees and alternatives allocations rise over 2001-2020); Sheets A and B adjust for them with fiscal-year-end fixed effects. "+DEF,
 [[('',2,2)],['CIOs','First-half mean','Second-half mean','Difference','Paired t-test p','Signed-rank p','Share increasing']],[34,8,13,13,12,12,12,12])
for gl,ys in GROUPS:
    ws.cell(row=r,column=1,value=gl).font=FB; r+=1
    for y in ys:
        p=PAIR[y]; ws.cell(row=r,column=1,value=LAB[y]).font=F
        for j,(v,fmt) in enumerate([(p['n'],'0'),(p['first'],'0.000'),(p['second'],'0.000'),(p['diff'],'0.000'),(p['p_t'],'0.000'),(p['p_sr'],'0.000'),(p['share_up'],'0.00')]):
            x=ws.cell(row=r,column=2+j,value=v); x.font=BLUE if j>=4 else F; x.alignment=RIGHT; x.number_format=fmt
        r+=1
for c in range(1,9): ws.cell(row=r-1,column=c).border=Border(bottom=thin)
wb.move_sheet('Estimates', offset=3)
out='/home/user/PensionData/referee_response/time_to_departure/halves_v3_fye1.xlsx'; wb.save(out); print('saved',out)
