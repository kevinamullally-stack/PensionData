import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
D=pd.read_pickle('returns_halves_equal_weight_fe.pkl'); R=D['R']; YS=D['YS']
TNR='Times New Roman'; SZ=10.5
F=Font(name=TNR,size=SZ); FB=Font(name=TNR,size=SZ,bold=True); FI=Font(name=TNR,size=SZ,italic=True)
thin=Side(style='thin'); CENTER=Alignment(horizontal='center',vertical='center',wrap_text=True); RIGHT=Alignment(horizontal='right',vertical='center'); LEFT=Alignment(horizontal='left',vertical='center',wrap_text=True); JUST=Alignment(horizontal='justify',wrap_text=True,vertical='top')
def st(p): return '' if p is None or np.isnan(p) else '***' if p<0.01 else '**' if p<0.05 else '*' if p<0.1 else ''
def cp(v,p): return f"{v:+.3f}{st(p)} ({p:.3f})"
W=[50,15,15,15,10,10,15,15,15,10,10]
ROWS=[('mean','Equal-weighted mean of spell changes (no FE)'),('mean_dm','  Net of the fiscal-year-end mean'),('fe','CIO FE + fiscal-year-end FE'),('fe_ctl','  CIO FE + fiscal-year-end FE + lagged plan controls'),
      ('fe_conf','Conflicted CIOs only: CIO FE + fiscal-year-end FE'),('fe_conf_ctl','  Conflicted CIOs only: + lagged plan controls')]
wb=Workbook(); ws=wb.active; ws.title='Equal-weighted halves'; ws.sheet_view.showGridLines=False
for j,w in enumerate(W): ws.column_dimensions[L(1+j)].width=w
ws['A1']='Change in Returns From the First to the Second Half of Tenure, Each CIO Weighted Equally: CIO and Fiscal-Year-End Fixed Effects, All Years and Excluding FY2009-2014'; ws['A1'].font=FB
note=("Sample: CIOs who have left (curr = 0) in the Table III sample with defined halves of tenure; the right-hand columns drop fiscal years 2009-2014 (the shock period). The second half is the years in which tenure exceeds half of tenure at departure, defined from each CIO's full tenure. A CIO-plan spell is conflicted if its CIO is a Conflicted CIO. "
      "The first two rows report the equal-weighted mean of spell-level changes (mean return in the second half minus the first half, every spell counted once), with t-test p-values (as in returns_halves_equal_weight.xlsx); 'net of the fiscal-year-end mean' subtracts the mean of the period's Table III plan-years with the same fiscal-year-end date. "
      "The regression rows estimate Return = CIO FE + fiscal-year-end FE + b x Second Half + d x Conflicted x Second Half, weighting each plan-year by one over the number of plan-years in its spell-half so that each spell's two halves carry equal weight; with CIO fixed effects alone this reproduces the equal-weighted mean exactly, and the fiscal-year-end fixed effects remove common shocks. "
      "CIO fixed effects are for each CIO's tenure at a plan. 'Other' is b (other departing CIOs), 'Conflicted' is b + d, and the difference is d. Lagged plan controls: investment return, log plan size, funding ratio, allocations, separate investment board, % political appointees (CIO tenure and age are omitted because tenure rises mechanically from the first to the second half). "
      "The conflicted-only rows estimate b on conflicted CIOs' spells alone. Entries are estimates with p-values in parentheses (standard errors double clustered by plan and year for the regressions); the 'plan p' column gives the p-value of the difference (conflicted change in the conflicted-only rows) with standard errors clustered by plan only. ***, **, and * denote significance at the 1%, 5%, and 10% level.")
ws.merge_cells('A2:K2'); ws['A2']=note; ws['A2'].font=F; ws['A2'].alignment=JUST; ws.row_dimensions[2].height=14*(len(note)//150+1)
r=4
for pi,(y,yl) in enumerate(YS):
    ws.cell(row=r,column=1,value=f"Panel {'AB'[pi]}. {yl}").font=FB; r+=1
    for c0,lab in [(2,'All fiscal years'),(7,'Excluding fiscal years 2009-2014')]:
        ws.merge_cells(start_row=r,start_column=c0,end_row=r,end_column=c0+4); x=ws.cell(row=r,column=c0,value=lab); x.font=FB; x.alignment=CENTER
        for cc in range(c0,c0+5): ws.cell(row=r,column=cc).border=Border(top=thin,bottom=thin)
    ws.cell(row=r,column=1).border=Border(top=thin); r+=1
    H=['','Conflicted: change (p)','Other departing CIOs: change (p)','Conflicted minus other (p)','Plan p','Spells (conflicted / other)']+['Conflicted: change (p)','Other departing CIOs: change (p)','Conflicted minus other (p)','Plan p','Spells (conflicted / other)']
    for j,h in enumerate(H): x=ws.cell(row=r,column=1+j,value=h); x.font=F; x.alignment=CENTER; x.border=Border(bottom=thin)
    ws.row_dimensions[r].height=44; r+=1
    for k,lab in ROWS:
        ws.cell(row=r,column=1,value=lab).font=FB if not lab.startswith('  ') else F; ws.cell(row=r,column=1).alignment=LEFT
        for c0,per in [(2,'all'),(7,'ex0914')]:
            q=R[(per,y,k)]
            if k.startswith('mean'):
                vals=[cp(q['mc'],q['p1c']),cp(q['mo'],q['p1o']),cp(q['diff'],q['p']),'',f"{q['nc']} / {q['no']}"]
            elif 'conf' in k:
                vals=[cp(q['cnf']['est'],q['cnf']['p']),'','',f"{q['cnf']['pp']:.3f}",f"{q['nsp_c']}"]
            else:
                vals=[cp(q['cnf']['est'],q['cnf']['p']),cp(q['oth']['est'],q['oth']['p']),cp(q['diff']['est'],q['diff']['p']),f"{q['diff']['pp']:.3f}",f"{q['nsp_c']} / {q['nsp_o']}"]
            for j,v in enumerate(vals): x=ws.cell(row=r,column=c0+j,value=v); x.font=FB if k=='fe' else F; x.alignment=RIGHT
        r+=1
    obs=[R[(per,y,'fe')]['N'] for per in ['all','ex0914']]
    ws.cell(row=r,column=1,value='Plan-years in the CIO FE + fiscal-year-end FE regression').font=FI
    for c0,n in zip([2,7],obs): x=ws.cell(row=r,column=c0,value=f"{n:,}"); x.font=FI; x.alignment=RIGHT
    r+=1
    for cc in range(1,12): ws.cell(row=r-1,column=cc).border=Border(bottom=thin)
    r+=2
ws.page_setup.orientation='landscape'; ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
out='/home/user/PensionData/referee_response/time_to_departure/returns_halves_equal_weight_fe.xlsx'; wb.save(out); print('saved',out)
