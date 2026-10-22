import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
O=pd.read_pickle('preshock_halves_power.pkl'); PERS=['full','pre2009','ex0914']; YS=[('peer_adj_ret100','Peer-Adjusted Return'),('b_AB','Benchmark-Adjusted Return')]
TNR='Times New Roman'; SZ=10.5
F=Font(name=TNR,size=SZ); FB=Font(name=TNR,size=SZ,bold=True); FI=Font(name=TNR,size=SZ,italic=True)
thin=Side(style='thin'); CENTER=Alignment(horizontal='center',vertical='center',wrap_text=True); RIGHT=Alignment(horizontal='right'); LEFT=Alignment(horizontal='left',vertical='center',wrap_text=True); JUST=Alignment(horizontal='justify',wrap_text=True,vertical='top')
wb=Workbook(); ws=wb.active; ws.title='Sample and power'; ws.sheet_view.showGridLines=False
for j,w in enumerate([62,20,20,20]): ws.column_dimensions[L(1+j)].width=w
ws['A1']='Sample Composition and Statistical Power of the First-Half vs. Second-Half Tests Outside the 2009-2014 Shock Period'; ws['A1'].font=FB
note=("Counts are for the estimation sample of column 1 of table3_conflicted_halves_preshock.xlsx (Table III controls, fiscal-year-end fixed effects); the full-sample column uses the same specification (table3_conflicted_halves_fye1.xlsx). "
      "Conflicted plan-years are split into the first and second half of the CIO's tenure at the plan, defined from full tenure. 'Observed in both halves' counts conflicted CIOs (CIO-plan tenures) with at least one estimation-sample plan-year in each half; only these identify the within-CIO comparison in column 4. "
      "Tenure at departure is the CIO's full tenure at the plan in years. The minimum detectable difference is the second-minus-first difference that a two-sided 5% test would detect with 80% power, (t(0.975) + t(0.80)) x standard error, with degrees of freedom equal to the number of year clusters minus one, using each column's double-clustered standard error. "
      "Columns: (1) fiscal-year-end FE; (2) plan FE; (3) plan + fiscal-year-end FE; (4) conflicted CIOs only, CIO + fiscal-year-end FE.")
ws.merge_cells('A2:D2'); ws['A2']=note; ws['A2'].font=F; ws['A2'].alignment=JUST; ws.row_dimensions[2].height=14*(len(note)//105+1)
r=4
for pi,(y,yl) in enumerate(YS):
    ws.cell(row=r,column=1,value=f"Panel {'AB'[pi]}. {yl}").font=FB; r+=1
    for j,h in enumerate(['','Full sample (FY2002-2020)','Through FY2008','Excluding FY2009-2014']):
        x=ws.cell(row=r,column=1+j,value=h); x.font=FB if j else F; x.alignment=CENTER; x.border=Border(top=thin,bottom=thin)
    ws.row_dimensions[r].height=30; r+=1
    o={p:O[(p,y)] for p in PERS}
    def row(lab,fn,bold=False,italic=False):
        global r
        ws.cell(row=r,column=1,value=lab).font=FB if bold else FI if italic else F; ws.cell(row=r,column=1).alignment=LEFT
        for j,p in enumerate(PERS): x=ws.cell(row=r,column=2+j,value=fn(o[p])); x.font=FB if bold else FI if italic else F; x.alignment=RIGHT
        r+=1
    def sub(lab):
        global r
        ws.cell(row=r,column=1,value=lab).font=FI; r+=1
    sub('Sample')
    row('  Plan-years',lambda q:f"{q['N']:,}")
    row('  Fiscal years (year clusters)',lambda q:f"{q['years']}")
    row('  Plans',lambda q:f"{q['plans']}")
    sub('Conflicted CIOs')
    row('  Conflicted plan-years',lambda q:f"{q['conf']} ({q['conf_share']:.0%})",bold=True)
    row('    in the first half of tenure',lambda q:f"{q['h1']}")
    row('    in the second half of tenure',lambda q:f"{q['h2']}")
    row('  Conflicted CIOs',lambda q:f"{q['conf_cios']}",bold=True)
    row('    CIO-plan tenures',lambda q:f"{q['conf_spells']}")
    row('    plans',lambda q:f"{q['conf_plans']}")
    row('  Conflicted CIOs observed in both halves',lambda q:f"{q['cios_both']}",bold=True)
    row('    CIO-plan tenures observed in both halves',lambda q:f"{q['spells_both']}")
    row('  Conflicted plan-years per CIO-plan tenure: mean (median)',lambda q:f"{q['py_per_spell']:.1f} ({q['py_per_spell_med']:.0f})")
    row('  Tenure at departure, years: mean (median)',lambda q:f"{q['ten_dep']:.1f} ({q['ten_dep_med']:.0f})")
    row('  Tenure in conflicted plan-years: first half / second half',lambda q:f"{q['ten_h1']:.1f} / {q['ten_h2']:.1f}")
    sub('Precision of the second-minus-first test')
    for j,lab in enumerate(['(1) FYE FE','(2) Plan FE','(3) Plan + FYE FE','(4) Conflicted only, CIO + FYE FE']):
        row(f'  {lab}: standard error / minimum detectable difference',lambda q,j=j:f"{q['se'][j]:.2f} / {q['mde'][j]:.2f}")
    row('  (4) Conflicted only: plan-years / CIO-plan tenures',lambda q:f"{q['n4']} / {q['sp4'][0]}")
    row('  Estimated second-minus-first difference, (1)-(3)',lambda q:' / '.join(f"{x:+.2f}" for x in q['diff'][:3]),italic=True)
    row('  For scale: Conflicted x First / x Second Half, column (1)',lambda q:f"{q['h1coef']:+.2f} / {q['h2coef']:+.2f}",italic=True)
    for c in range(1,5): ws.cell(row=r-1,column=c).border=Border(bottom=thin)
    r+=2
ws.page_setup.orientation='portrait'; ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
out='/home/user/PensionData/referee_response/time_to_departure/preshock_halves_power.xlsx'; wb.save(out); print('saved',out)
