import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
D=pd.read_pickle('ret_halves_ew.pkl'); R=D['R']; SL=D['SL']; YS=D['YS']
TNR='Times New Roman'; SZ=10.5
F=Font(name=TNR,size=SZ); FB=Font(name=TNR,size=SZ,bold=True); FI=Font(name=TNR,size=SZ,italic=True)
thin=Side(style='thin'); CENTER=Alignment(horizontal='center',vertical='center',wrap_text=True); RIGHT=Alignment(horizontal='right',vertical='center'); LEFT=Alignment(horizontal='left',vertical='center',wrap_text=True); JUST=Alignment(horizontal='justify',wrap_text=True,vertical='top')
def st(p): return '' if p is None or np.isnan(p) else '***' if p<0.01 else '**' if p<0.05 else '*' if p<0.1 else ''
def cp(v,p,dec=3): return f"{v:+.{dec}f}{st(p)} ({p:.3f})"
wb=Workbook(); ws=wb.active; ws.title='Returns, equal weight'; ws.sheet_view.showGridLines=False
W=[44,18,20,18,18,10,10,18,12]
for j,w in enumerate(W): ws.column_dimensions[L(1+j)].width=w
ws['A1']='Change in Returns From the First to the Second Half of Tenure, Each CIO Weighted Equally'; ws['A1'].font=FB
note=("Sample: CIOs who have left (curr = 0) in the Table III sample with defined halves of tenure. For each CIO-plan spell (or each CIO, averaging across plans within a CIO-year), the change is the mean return in the second half of tenure minus the mean return in the first half; every spell (CIO) counts once. "
      "The second half is the years in which tenure exceeds half of tenure at departure. Columns report the mean change for conflicted CIOs with the p-value of a test that it is zero, their median change with the signed-rank p-value, the mean change for other departing CIOs, the conflicted-minus-other difference with p-values from a t-test, a t-test allowing unequal variances (Welch), and a rank-sum test, and the difference when spells are weighted by their plan-years (standard errors clustered by plan). "
      "'Net of fiscal-year-end mean' subtracts the mean return of all Table III plan-years with the same fiscal-year-end date. 'Drop' rows remove the conflicted spells with the most negative or most positive changes. The last row of each panel repeats the comparison with each spell's OLS slope of return on fiscal year (at least three years, winsorized at 1% and 99%), parallel to the Fee Slope. "
      "Entries are estimates with p-values in parentheses; ***, **, and * denote significance at the 1%, 5%, and 10% level.")
ws.merge_cells('A2:I2'); ws['A2']=note; ws['A2'].font=F; ws['A2'].alignment=JUST; ws.row_dimensions[2].height=14*(len(note)//125+2)
r=4
HEAD=['','Conflicted: mean change (p vs. 0)','Conflicted: median change (signed-rank p)','Other departing CIOs: mean change (p vs. 0)','Conflicted minus other (t-test p)','Welch p','Rank-sum p','Weighted by plan-years (p)','Spells (conflicted / other)']
ROWS=[('spell','base','Each CIO-plan spell'),('spell','ten4','  Tenure at departure of at least 4 years'),('spell','dm','  Net of fiscal-year-end mean'),
      ('spell','dropneg3','  Drop the 3 most negative conflicted changes'),('spell','droppos3','  Drop the 3 most positive conflicted changes'),
      ('spell','dropneg5','  Drop the 5 most negative conflicted changes'),('spell','droppos5','  Drop the 5 most positive conflicted changes'),
      ('cio','base','One observation per CIO'),('cio','ten4','  Tenure at departure of at least 4 years'),('cio','dm','  Net of fiscal-year-end mean')]
for y,yl in YS:
    ws.cell(row=r,column=1,value=f"Panel {'A' if y.startswith('peer') else 'B'}. {yl}").font=FB; r+=1
    for j,h in enumerate(HEAD): x=ws.cell(row=r,column=1+j,value=h); x.font=F; x.alignment=CENTER; x.border=Border(top=thin,bottom=thin)
    ws.row_dimensions[r].height=58; r+=1
    for unit,k,lab in ROWS:
        q=R[(y,unit,k)]
        ws.cell(row=r,column=1,value=lab).font=FB if not lab.startswith('  ') else F; ws.cell(row=r,column=1).alignment=LEFT
        vals=[cp(q['mc'],q['p1c']),cp(q['medc'],q['psr']),cp(q['mo'],q['p1o']),cp(q['diff'],q['p']),f"{q['pw']:.3f}",f"{q['prs']:.3f}",cp(q['wdiff'],q['wp']),f"{q['nc']} / {q['no']}"]
        for j,v in enumerate(vals): x=ws.cell(row=r,column=2+j,value=v); x.font=F; x.alignment=RIGHT
        r+=1
    q=SL[y]; ws.cell(row=r,column=1,value='Spell-level return slope (pp per year)').font=FI
    vals=[cp(q['mc'],q['p1c'],4),f"{q['medc']:+.4f}",cp(q['mo'],q['p1o'],4),cp(q['diff'],q['p'],4),f"{q['pw']:.3f}",f"{q['prs']:.3f}",cp(q['wdiff'],q['wp'],4),f"{q['nc']} / {q['no']}"]
    for j,v in enumerate(vals): x=ws.cell(row=r,column=2+j,value=v); x.font=FI; x.alignment=RIGHT
    r+=1
    for c in range(1,10): ws.cell(row=r-1,column=c).border=Border(bottom=thin)
    r+=1
ws.page_setup.orientation='landscape'; ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
out='/home/user/PensionData/referee_response/time_to_departure/returns_halves_equal_weight.xlsx'; wb.save(out); print('saved',out)
